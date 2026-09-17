"""Clean Excel presentation with normalized SEC source and lineage sheets.

The financial engine owns selection and calculation.  This module only renders
the final pivot and the ``fact_audit`` ledger attached to ``DataFrame.attrs``.
It never fetches data or changes a financial value.
"""

from __future__ import annotations

import hashlib
import json
import math
import platform
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

SHEET_NAMES = (
    "00_OVERVIEW",
    "01_ANNUAL",
    "02_QUARTERLY",
    "03_INCOME_STATEMENT",
    "04_BALANCE_SHEET",
    "05_CASH_FLOW",
    "06_SEGMENTS",
    "07_OPERATING_KPI",
    "08_QA",
    "09_REVIEW",
    "90_SOURCES",
    "91_LINEAGE",
    "97_RAW_KPI",
    "98_RAW_QA",
)

_HEADER = PatternFill("solid", fgColor="17365D")
_SECTION = PatternFill("solid", fgColor="D9EAF7")
_SUBTLE = PatternFill("solid", fgColor="F3F6F9")
_WHITE_BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_TITLE = Font(name="Arial", size=18, bold=True, color="17365D")
_BODY = Font(name="Arial", size=10, color="222222")
_BOLD = Font(name="Arial", size=10, bold=True, color="222222")
_LINK = Font(name="Arial", size=10, color="374151", underline=None)
_THIN_BLUE = Side(style="thin", color="9EBDD5")
_TOTAL_BORDER = Border(top=_THIN_BLUE, bottom=_THIN_BLUE)
_ALIGN_LEFT = Alignment(horizontal="left")
_ALIGN_RIGHT = Alignment(horizontal="right")
_ALIGN_TOP = Alignment(horizontal="left", vertical="top")
_ALIGN_WRAP_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

_MAIN_CATEGORIES = {
    "03_INCOME_STATEMENT": ("1_Income_Statement",),
    "04_BALANCE_SHEET": ("2_Balance_Sheet",),
    "05_CASH_FLOW": ("3_Cash_Flow",),
    "06_SEGMENTS": (
        "4a_Segments_Business",
        "4a_Segments_Platform",
        "4b_Segments_Product_Type",
        "4b_Segments_Geographic_Regions",
        "4c_Segments_Geographic_Regions",
        "4c_Segments_Geographic_Countries",
        "4d_Segments_Cross_Tabulated",
        "4e_Operating_Metrics",
    ),
}

_RAW_KPI_CATEGORIES = ("5_KPI_Metrics", "6_Disclosures", "7_Concentration_Risk")
_RAW_QA_CATEGORIES = ("8_Integrity_Checks",)

_KEY_METRICS = (
    ("Revenue", ("Revenue",)),
    ("Operating Income", ("Operating Income",)),
    ("Net Income", ("Net Income",)),
    ("Operating Cash Flow", ("Operating Cash Flow",)),
    ("Cash & Equivalents", ("Cash & Equivalents", "Cash and Cash Equivalents")),
    ("Total Debt", ("Metric: Total Debt", "Total Debt")),
)

_TLN_KPI_SPECS = (
    ("PJM Adjusted EBITDA", "Adjusted Segment EBITDA - PJM", "Adjusted Segment EBITDA - PJM Segment Adjusted EBITDA"),
    ("PJM Revenue", "Revenue - PJM"),
    ("PJM Capital Expenditures", "Capital Expenditures - PJM"),
    ("PJM West Hub ATC", "Other Segment Items - $/MWh - PJM West Hub ATC"),
    ("PJM West Hub ATC Spark Spread", "Other Segment Items - $/MWh (a) - PJM West Hub ATC Spark Spreads"),
    ("Standardized Free Cash Flow", "Metric: Free Cash Flow", "Free Cash Flow"),
    ("Total Debt", "Metric: Total Debt"),
    ("Net Debt", "Metric: Net Cash (Debt)"),
    ("Operating Margin", "Operating Margin (%)"),
    ("Net Margin", "Net Margin (%)"),
)

_INDUSTRY_COVERAGE = {
    "TLN": "POWER / SUPPORTED",
}

_EQUITY_ISSUANCE_CASH_LABELS = {
    "Shares Issued",
    "Shares Issued - Common Stock",
    "Shares Issued - Preferred Stock",
    "Shares Issued - Mandatory Convertible Preferred Stock",
    "Shares Issued - Combined Common and Preferred Stock",
    "Shares Issued - Other Equity",
    "Shares Issued (Stock Plans)",
    "Net Shares Issued (Repurchased)",
}

_FRONT_LABELS = {
    "Net Cash Flow": "Net Cash Flow before FX",
    "Shares Issued": "Equity Issuance Proceeds",
    "Shares Issued - Common Stock": "Proceeds from Common Stock Issuance",
    "Shares Issued - Preferred Stock": "Proceeds from Preferred Stock Issuance",
    "Shares Issued - Mandatory Convertible Preferred Stock": (
        "Proceeds from Mandatory Convertible Preferred Stock Issuance"
    ),
    "Shares Issued - Combined Common and Preferred Stock": (
        "Proceeds from Common and Preferred Stock Issuance"
    ),
    "Shares Issued - Other Equity": "Proceeds from Other Equity Issuance",
    "Shares Issued (Stock Plans)": "Proceeds from Stock Plans",
    "Net Shares Issued (Repurchased)": "Net Equity Issuance (Repurchase) Cash Flow",
}

_LINEAGE_AUDIT_FIELDS = (
    "Ticker", "Category", "Label", "Period", "Value", "Accession", "FilingUrl",
    "Concept", "Form", "Filed", "AcceptedAt", "Accepted", "AcceptanceDateTime",
    "FY", "Q", "Start", "End", "SourceKind", "SourceAdmissionRule",
    "SourceDerivation", "SourceDerivationFormula", "SourceDerivationExpression",
    "SourceSemanticType", "SourceAnalyticalBasis",
    "SourceClassificationConfidence",
    "SourceDimensionAxes", "SourceDimensionMembers", "SourceUnitKind",
    "SourceUnitScale", "SourceRawValue", "SourceReportedValue", "SourceCellIdentity", "SourceTableFingerprint",
    "SourceDisplaySign", "SourceInputAccessions", "SourceInputFacts",
    "SourceExactUnitKind", "SourceEvidenceAxes", "SourceEvidenceMembers", "SourceConceptQName", "SourceTableEvidence", "SourceEvidenceStatus",
)

# Engine reconciliation residuals are diagnostics, never reported face lines.
_ENGINE_DIAGNOSTIC_LABELS = frozenset({
    "Other Operating Adjustments (Net)", "Other Investing Adjustments (Net)",
    "Other Financing Adjustments (Net)", "Operating Income: Other Adjustments",
})

_SOURCE_EXPORT_FIELDS = (
    "Source ID", "Ticker", "Form", "Report Period", "Filed", "Accepted At",
    "Accession", "SEC Link", "Evidence Class",
)
_LINEAGE_EXPORT_FIELDS = (
    "Lineage ID", "Excel Sheet", "Metric", "Period", "Value",
    "Raw Reported Value", "Type", "Source ID", "Input Source IDs",
    "XBRL Concept", "Context", "Unit", "Unit Scale", "Display Scale",
    "Display Unit", "Formula", "Input Facts", "Method", "Locator",
    "Semantic Classification", "Semantic Confidence", "Review Required", "Table Evidence",
)
_QA_EXPORT_FIELDS = (
    "Check ID", "Status", "Severity", "Check", "Failures", "Evidence",
)


def _present(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return math.isfinite(float(value))
    try:
        return not bool(pd.isna(value))
    except (TypeError, ValueError):
        return True


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if not _present(value):
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def _safe_text(value: Any) -> str:
    text = _text(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _number(value: Any) -> float | int | None:
    if not _present(value) or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        parsed = float(value)
        return int(parsed) if parsed.is_integer() else parsed
    try:
        parsed = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    if not math.isfinite(parsed):
        return None
    return int(parsed) if parsed.is_integer() else parsed


def _same_number(left: Any, right: Any) -> bool:
    a, b = _number(left), _number(right)
    if a is None or b is None:
        return _text(left) == _text(right)
    return math.isclose(float(a), float(b), rel_tol=1e-12, abs_tol=1e-9)


def _decimal(value: Any) -> Decimal | None:
    """Parse a finite value without introducing another float round trip."""
    if not _present(value) or isinstance(value, bool):
        return None
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return parsed if parsed.is_finite() else None


def _period_key(period: Any, end_date: Any = None) -> tuple[Any, ...]:
    if _present(end_date):
        parsed = pd.to_datetime(end_date, errors="coerce")
        if pd.notna(parsed):
            return (0, int(parsed.year), int(parsed.month), int(parsed.day), _text(period))
    label = _text(period)
    match = re.match(r"^(\d{4})\s*[-_/ ]?\s*Q([1-4])$", label, re.I)
    if match:
        year, quarter = int(match.group(1)), int(match.group(2))
        return (0, year, quarter * 3, 31, label)
    match = re.match(r"^(?:FY\s*)?(\d{4})(?:\s*[-_/ ]?\s*FY)?$", label, re.I)
    if match:
        return (0, int(match.group(1)), 12, 31, label)
    return (1, label)


def _quarter_ordinal(period: Any) -> int | None:
    match = re.fullmatch(r"(\d{4})-Q([1-4])", _text(period), re.I)
    if not match:
        return None
    return int(match.group(1)) * 4 + int(match.group(2))


def _canonical_categories(frame: pd.DataFrame) -> list[str]:
    if not isinstance(frame.index, pd.MultiIndex) or frame.index.nlevels < 2:
        raise ValueError("final_pivot must use a (Category, Label) MultiIndex")
    return list(dict.fromkeys(str(x) for x in frame.index.get_level_values(0)))


def _periods(frame: pd.DataFrame) -> tuple[list[Any], Mapping[Any, Any]]:
    dates: dict[Any, Any] = dict(frame.attrs.get("period_dates") or {})
    key = ("0_Period_Header", "Period Ending")
    if key in frame.index:
        row = frame.loc[key]
        if isinstance(row, pd.DataFrame):
            row = row.iloc[0]
        dates = {column: row.get(column) for column in frame.columns}
    if not dates:
        audit = frame.attrs.get("fact_audit")
        if isinstance(audit, pd.DataFrame) and {"Period", "End"}.issubset(audit.columns):
            for period, group in audit.groupby("Period", dropna=False):
                ends = sorted({_text(value) for value in group["End"] if _text(value) not in {"", "None", "nan"}})
                if len(ends) == 1:
                    dates[period] = ends[0]
    original = list(frame.columns)
    ordered = sorted(original, key=lambda item: _period_key(item, dates.get(item)))
    return ordered, dates


def _row_dict(row: Any) -> dict[str, Any]:
    if isinstance(row, pd.Series):
        return row.to_dict()
    if isinstance(row, Mapping):
        return dict(row)
    return {}


def _audit_records(audit: pd.DataFrame, fields: Iterable[str]) -> list[dict[str, Any]]:
    columns = [field for field in fields if field in audit.columns]
    return audit.loc[:, columns].to_dict(orient="records") if columns else []


def _audit_groups(audit: pd.DataFrame) -> dict[tuple[str, str, str], list[dict[str, Any]]]:
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    if audit is None or audit.empty:
        return groups
    required = {"Category", "Label", "Period", "Value"}
    if not required.issubset(audit.columns):
        return groups
    for item in _audit_records(audit, _LINEAGE_AUDIT_FIELDS):
        key = (_text(item.get("Category")), _text(item.get("Label")), _text(item.get("Period")))
        groups[key].append(item)
    return groups


def _source_completeness(row: Mapping[str, Any]) -> tuple[int, str, str]:
    fields = ("Accession", "FilingUrl", "Concept", "Form", "Filed", "Start", "End")
    score = sum(bool(_text(row.get(field))) for field in fields)
    return (-score, _text(row.get("Filed")), _text(row.get("Accession")))


def _match_audit(
    groups: Mapping[tuple[str, str, str], list[dict[str, Any]]],
    category: str,
    label: str,
    period: Any,
    value: Any,
) -> dict[str, Any]:
    candidates = groups.get((category, label, _text(period)), [])
    exact = [row for row in candidates if _same_number(row.get("Value"), value)]
    pool = exact or candidates
    return sorted(pool, key=_source_completeness)[0] if pool else {}


def _origin(category: str, row: Mapping[str, Any]) -> str:
    kind = _text(row.get("SourceKind")).casefold()
    derivation = _text(row.get("SourceDerivation") or row.get("SourceDerivationFormula"))
    if "manual" in kind:
        return "M"
    if "calculat" in kind:
        return "C"
    if derivation:
        return "D"
    if "calculat" in kind or category.startswith("8_"):
        return "C"
    if any(token in kind for token in ("xbrl", "html", "table", "sec", "report", "filing", "disclosure")):
        return "R"
    if category.startswith("5_"):
        return "C"
    return ""


def _unit(label: str, row: Mapping[str, Any]) -> str:
    exact = _text(row.get('SourceExactUnitKind'))
    if exact:
        return {'USD/shares': 'USD/share'}.get(exact, exact)
    lowered = label.casefold()
    if label in _EQUITY_ISSUANCE_CASH_LABELS:
        return "USD"
    if "/mwh" in lowered or "per mwh" in lowered:
        return "USD/MWh"
    if "twh" in lowered:
        return "TWh"
    if re.search(r"\bmwh\b", lowered):
        return "MWh"
    if re.search(r"\bmw\b", lowered):
        return "MW"
    if "per share" in lowered or "eps" in lowered:
        return "USD/share"
    if re.search(r'\bshares? (?:outstanding|count)\b|\bweighted.average shares?\b', lowered):
        return "shares"
    kind = _text(row.get("SourceUnitKind")).casefold()
    if kind in {"money", "currency", "usd"}:
        return "USD"
    if "rate securities" in lowered:
        return "USD"
    if "%" in lowered or "percent" in lowered or "margin" in lowered or re.search(
        r"\b(?:rate|yield|factor)\b", lowered
    ):
        return "%"
    if any(token in lowered for token in (
        "revenue", "income", "expense", "cash flow", "free cash", "ebitda",
        "debt", "expenditure", "capex", "assets", "liabilities", "equity",
    )):
        return "USD"
    if kind:
        return _text(row.get("SourceUnitKind"))
    category = _text(row.get("Category"))
    if category.startswith((
        "4a_Segments_Business",
        "4a_Segments_Platform",
        "4b_Segments_Geographic",
        "4c_Segments_Geographic",
        "4d_Segments_Cross_Tabulated",
    )):
        return "USD"
    if category_is_financial(_text(row.get("Category"))):
        return "USD"
    return ""


def category_is_financial(category: str) -> bool:
    return category.startswith(("1_", "2_", "3_"))


def _context(row: Mapping[str, Any]) -> str:
    start, end = _text(row.get("Start")), _text(row.get("End"))
    period = end if not start or start == end else f"{start}..{end}"
    axes = _text(row.get('SourceEvidenceAxes')) or _text(row.get("SourceDimensionAxes"))
    members = _text(row.get('SourceEvidenceMembers')) or _text(row.get("SourceDimensionMembers"))
    dimensions = "; ".join(part for part in (axes, members) if part)
    return " | ".join(part for part in (period, dimensions) if part)


def _formula(row: Mapping[str, Any]) -> str:
    return _text(
        row.get("SourceDerivationFormula")
        or row.get("SourceDerivationExpression")
        or row.get("SourceDerivation")
    )


def _input_accessions(row: Mapping[str, Any]) -> list[str]:
    accessions: list[str] = []
    for key, value in row.items():
        if "accession" not in str(key).casefold():
            continue
        for token in re.split(r"[;,|\s]+", _text(value)):
            if re.fullmatch(r"\d{10}-\d{2}-\d{6}", token) and token not in accessions:
                accessions.append(token)
    return accessions


def _source_rows(audit: pd.DataFrame) -> tuple[list[dict[str, str]], dict[str, str]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    if audit is not None and not audit.empty:
        for row in _audit_records(audit, _LINEAGE_AUDIT_FIELDS):
            accession = _text(row.get("Accession"))
            if accession:
                grouped[accession].append(row)

    def first(rows: list[dict[str, Any]], field: str) -> str:
        return next((_text(row.get(field)) for row in rows if _text(row.get(field))), "")

    records: list[dict[str, str]] = []
    for accession, rows in grouped.items():
        form = first(rows, "Form")
        dated = []
        for row in rows:
            parsed = pd.to_datetime(row.get("End"), errors="coerce")
            if pd.notna(parsed):
                dated.append((parsed, row))
        latest = max(dated, key=lambda item: item[0])[1] if dated else rows[0]
        year = _text(latest.get("FY")) or (
            str(max(item[0] for item in dated).year) if dated else ""
        )
        quarter = _text(latest.get("Q"))
        report_period = f"FY{year}" if "10-K" in form.upper() else " ".join(part for part in (year, quarter) if part)
        records.append({
            "Ticker": first(rows, "Ticker"),
            "Form": form,
            "Report Period": report_period,
            "Filed": first(rows, "Filed"),
            "Accepted At": (
                first(rows, "AcceptedAt")
                or first(rows, "Accepted")
                or first(rows, "AcceptanceDateTime")
            ),
            "Accession": accession,
            "SEC Link": first(rows, "FilingUrl"),
            "Evidence Class": "SEC filing",
        })
    ordered = sorted(records, key=lambda row: (row["Filed"], row["Accession"]))
    accession_to_id: dict[str, str] = {}
    for number, row in enumerate(ordered, 1):
        source_id = f"F{number:03d}"
        row["Source ID"] = source_id
        accession_to_id[row["Accession"]] = source_id
    return ordered, accession_to_id


def _target_sheet(category: str) -> str:
    for sheet, categories in _MAIN_CATEGORIES.items():
        if category in categories:
            return sheet
    if category in _RAW_KPI_CATEGORIES:
        return "97_RAW_KPI"
    if category in _RAW_QA_CATEGORIES:
        return "98_RAW_QA"
    return "91_LINEAGE"


def _display_transform(label: str, unit: str, row: Mapping[str, Any] | None = None) -> tuple[float, str]:
    """Return presentation-only multiplier and readable unit."""
    if unit == "USD":
        sign = _number((row or {}).get("SourceDisplaySign"))
        return 1e-6 * (float(sign) if sign in {-1, 1} else 1.0), "$mm"
    if unit == "shares":
        return 1e-6, "mm shares"
    if unit == "USD/MWh":
        scale = _number((row or {}).get("SourceUnitScale"))
        return (1.0 / float(scale), "USD/MWh") if scale and scale > 1 else (1.0, "USD/MWh")
    return 1.0, unit


def _build_lineage(
    frame: pd.DataFrame,
    audit: pd.DataFrame,
    periods: Iterable[Any],
    accession_to_id: Mapping[str, str],
) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    # pandas propagates/deep-copies the large audit ledger into each Series.
    # Presentation needs the matrix only; source audit is passed separately.
    frame = pd.DataFrame(frame.to_numpy(copy=False), index=frame.index, columns=frame.columns)
    groups = _audit_groups(audit)
    period_rows: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    period_rows_any: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for (category, _label, period), candidates in groups.items():
        period_rows[(category, period)].extend(candidates)
        period_rows_any[period].extend(candidates)
    lineage: list[dict[str, Any]] = []
    issues: list[dict[str, str]] = []
    suppressed_noncore = 0
    for (category, label), values in frame.iterrows():
        category, label = _text(category), _text(label)
        if category == "0_Period_Header":
            continue
        for period in periods:
            value = values.get(period)
            if not _present(value):
                continue
            row = _match_audit(groups, category, label, period, value)
            accession = _text(row.get("Accession"))
            source_id = accession_to_id.get(accession, "")
            inputs = [accession_to_id[a] for a in _input_accessions(row) if a in accession_to_id]
            origin = _origin(category, row)
            period_evidence = period_rows.get((category, _text(period)), []) or period_rows_any.get(_text(period), [])
            if not source_id and period_evidence:
                fallback_accessions = list(dict.fromkeys(
                    _text(candidate.get("Accession")) for candidate in period_evidence
                    if _text(candidate.get("Accession")) in accession_to_id
                ))
                inputs = [accession_to_id[item] for item in fallback_accessions]
                source_id = inputs[0] if inputs else ""
                if not row:
                    origin = "C"
            if not source_id and origin == "D" and inputs:
                source_id = inputs[0]
            unit = _unit(label, {**row, "Category": category})
            display_scale, display_unit = _display_transform(label, unit, row)
            raw_confidence = _number(row.get("SourceClassificationConfidence"))
            if not row:
                semantic_confidence = "LOW"
            elif raw_confidence is not None:
                semantic_confidence = (
                    "HIGH" if float(raw_confidence) >= 0.95
                    else "REVIEW" if float(raw_confidence) >= 0.75
                    else "LOW"
                )
            elif category.startswith(("1_", "2_", "3_")):
                semantic_confidence = "HIGH"
            else:
                semantic_confidence = "LOW"
            semantic_classification = _text(row.get("SourceSemanticType"))
            if not semantic_classification:
                semantic_classification = (
                    "Standard financial metric"
                    if category.startswith(("1_", "2_", "3_"))
                    else "Unclassified metric"
                )
            lineage_id = f"L{len(lineage) + 1:06d}"
            formula = _formula(row)
            review_required = (
                semantic_confidence != "HIGH"
                or semantic_classification == "Unclassified metric"
                or not source_id
                or (origin in {"C", "D"} and not formula)
            )
            item = {
                "_Category": category,
                "Lineage ID": lineage_id,
                "Excel Sheet": "91_LINEAGE" if label in _ENGINE_DIAGNOSTIC_LABELS else _target_sheet(category),
                "Metric": label,
                "Period": _text(period),
                "Value": value,
                "Raw Reported Value": _text(row.get("SourceRawValue")),
                "Type": origin,
                "Source ID": source_id,
                "Input Source IDs": ", ".join(dict.fromkeys(inputs)),
                "XBRL Concept": _text(row.get('SourceConceptQName')) or _text(row.get("Concept")),
                "Context": _context(row),
                "Unit": unit,
                "Unit Scale": _number(row.get("SourceUnitScale")),
                "Display Scale": display_scale,
                "Display Unit": display_unit,
                "Formula": formula,
                "Input Facts": _text(row.get("SourceInputFacts")),
                "Method": _text(row.get("SourceAdmissionRule") or row.get("SourceKind")) or (
                    "engine calculation; period-level input filings" if period_evidence else ""
                ),
                "Locator": _text(row.get("SourceCellIdentity") or row.get("SourceTableFingerprint")),
                "Semantic Classification": semantic_classification,
                "Semantic Confidence": semantic_confidence,
                "Review Required": "YES" if review_required else "NO",
                "Table Evidence": _text(row.get('SourceTableEvidence')),
            }
            lineage.append(item)
            review_required = category.startswith(("1_", "2_", "3_", "4", "5_"))
            if not row and review_required and not source_id:
                issues.append(_issue("HIGH", label, period, "", "Missing lineage", "No matching fact_audit row or period-level input filing for the displayed value."))
            elif not row and review_required:
                issues.append(_issue(
                    "WARNING", label, period, source_id, "Calculated lineage granularity",
                    "The engine output is linked to period-level input filings; no exact selected-fact row was retained.",
                ))
            elif not row:
                suppressed_noncore += 1
            elif origin in {"R", "D"} and not source_id and review_required:
                issues.append(_issue("HIGH", label, period, "", "Missing source", "Reported or derived value has no filing accession."))
            if origin in {"C", "D"} and not item["Formula"]:
                issues.append(_issue("HIGH", label, period, source_id, "Missing formula", "Calculated or derived value has no retained expression."))
            if review_required and _number(value) is not None and not unit:
                issues.append(_issue("MEDIUM", label, period, source_id, "Missing unit", "Review the unit before relying on this value."))
    if suppressed_noncore:
        issues.append(_issue(
            "WARNING", "Non-core disclosures", "Multiple", "",
            "Non-core lineage coverage",
            f"{suppressed_noncore} raw disclosure or QA cells have no direct fact_audit match.",
        ))
    return lineage, issues


def _issue(severity: str, metric: Any, period: Any, source_id: str, check: str, finding: str) -> dict[str, str]:
    return {
        "Status": "OPEN",
        "Severity": severity,
        "Metric": _text(metric),
        "Period": _text(period),
        "Source ID": source_id,
        "Check": check,
        "Finding": finding,
        "Suggested Action": "Inspect 91_LINEAGE and the linked SEC filing.",
    }


def _write_title(ws: Any, title: str, subtitle: str = "") -> None:
    ws["B2"] = _safe_text(title)
    ws["B2"].font = _TITLE
    ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 42
    if subtitle:
        ws["B3"] = _safe_text(subtitle)
        ws["B3"].font = Font(name="Arial", size=10, color="666666", italic=True)
        ws["B3"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[3].height = 48
    ws.sheet_view.showGridLines = False


def _write_header(ws: Any, row: int, headers: list[str], start_col: int = 2) -> None:
    for offset, header in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + offset, value=header)
        cell.fill = _HEADER
        cell.font = _WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_value(cell: Cell, value: Any, label: str = "", multiplier: float = 1.0) -> None:
    numeric = _number(value)
    if numeric is None:
        cell.value = _safe_text(value)
        cell.alignment = _ALIGN_LEFT
        return
    cell.value = float(numeric) * multiplier
    lowered = label.casefold()
    cell.number_format = "#,##0.00;(#,##0.00);0" if any(
        token in lowered for token in ("per share", "eps", "margin", "ratio", "%", "factor")
    ) else "#,##0;(#,##0);0"
    cell.alignment = _ALIGN_RIGHT


def _front_value_is_eligible(item: Mapping[str, Any]) -> bool:
    """Keep unresolved values auditable without presenting them as trusted facts."""
    if _text(item.get("Review Required")) != "NO":
        return False
    origin = _text(item.get("Type"))
    if origin in {"C", "D"}:
        return bool(
            _text(item.get("Formula"))
            and (_text(item.get("Source ID")) or _text(item.get("Input Source IDs")))
        )
    if origin == "R":
        return bool(
            _text(item.get("Source ID"))
            and (_text(item.get("XBRL Concept")) or _text(item.get("Locator")))
        )
    return origin == "M"


def _write_front_value(
    cell: Cell,
    value: Any,
    label: str,
    item: Mapping[str, Any],
) -> None:
    if _present(value) and _front_value_is_eligible(item):
        multiplier = float(_number(item.get("Display Scale")) or 1.0)
        _write_value(cell, value, label, multiplier)
        return
    cell.value = "—" if _present(value) else ""
    cell.alignment = _ALIGN_RIGHT
    if _present(value):
        cell.font = Font(name="Arial", size=10, italic=True, color="9A6700")


def _set_internal_link(cell: Cell, row: int) -> None:
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'91_LINEAGE'!A{row}")
    cell.font = _LINK


def _is_total(label: str) -> bool:
    return label.startswith("Total ") or label in {
        "Revenue", "Gross Profit", "Operating Income", "Pretax Income", "Net Income",
        "Operating Cash Flow", "Investing Cash Flow", "Financing Cash Flow", "Net Cash Flow",
    }


def _write_matrix(
    ws: Any,
    frame: pd.DataFrame,
    periods: list[Any],
    dates: Mapping[Any, Any],
    categories: Iterable[str],
    lineage_rows: Mapping[tuple[str, str, str], int],
    lineage_items: Mapping[tuple[str, str, str], Mapping[str, Any]],
    title: str,
    hide_disclosures: bool = False,
) -> None:
    frame = pd.DataFrame(frame.to_numpy(copy=False), index=frame.index, columns=frame.columns)
    _write_title(ws, title, "USD / shares: millions; per-share / operating metrics: original units.")
    headers = ["Line Item", *[_text(period) for period in periods]]
    _write_header(ws, 5, headers)
    for offset, period in enumerate(periods, 3):
        ws.cell(6, offset, _safe_text(dates.get(period)))
        ws.cell(6, offset).font = Font(name="Arial", size=9, italic=True, color="666666")
        ws.cell(6, offset).alignment = Alignment(horizontal="center")
    ws.cell(6, 2, "Period Ending").font = _BOLD
    output_row = 7
    present_categories = set(_canonical_categories(frame))
    for category in categories:
        if category not in present_categories:
            continue
        category_rows = frame.xs(category, level=0, drop_level=True)
        ws.cell(output_row, 2, re.sub(r"^\d+[a-z]?_", "", category).replace("_", " "))
        ws.cell(output_row, 2).fill = _SECTION
        ws.cell(output_row, 2).font = _BOLD
        for col in range(3, len(periods) + 3):
            ws.cell(output_row, col).fill = _SECTION
        output_row += 1
        for label, values in category_rows.iterrows():
            label_text = _text(label)
            if label_text in _ENGINE_DIAGNOSTIC_LABELS:
                continue
            front_label = _FRONT_LABELS.get(label_text, label_text)
            label_cell = ws.cell(output_row, 2, front_label)
            label_cell.font = _BOLD if _is_total(label_text) else _BODY
            label_cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                indent=0 if _is_total(label_text) else 1,
                wrap_text=len(front_label) > 42,
            )
            if len(front_label) > 42:
                ws.row_dimensions[output_row].height = 28
            if _is_total(label_text):
                label_cell.border = _TOTAL_BORDER
            for offset, period in enumerate(periods, 3):
                value = values.get(period)
                cell = ws.cell(output_row, offset)
                key = (category, label_text, _text(period))
                lineage_item = lineage_items.get(key, {})
                _write_front_value(cell, value, label_text, lineage_item)
                target = lineage_rows.get(key)
                if target and _present(value):
                    _set_internal_link(cell, target)
                if _is_total(label_text):
                    cell.border = _TOTAL_BORDER
            if hide_disclosures and category == "6_Disclosures":
                ws.row_dimensions[output_row].hidden = True
                ws.row_dimensions[output_row].outlineLevel = 1
            output_row += 1
    ws.freeze_panes = "C7"
    ws.auto_filter.ref = f"B5:{get_column_letter(len(periods) + 2)}{max(output_row - 1, 6)}"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 44
    for col in range(3, len(periods) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 2
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:6"
    ws.print_title_cols = "B:B"
    ws.print_area = f"B2:{get_column_letter(len(periods) + 2)}{max(output_row - 1, 6)}"


def _summary_rows(frame: pd.DataFrame) -> list[tuple[str, str, str]]:
    by_label: dict[str, list[str]] = defaultdict(list)
    for category, label in frame.index:
        if _text(category) != "0_Period_Header":
            by_label[_text(label)].append(_text(category))
    selected: list[tuple[str, str, str]] = []
    for display, aliases in _KEY_METRICS:
        exact = [
            (label, cats) for label, cats in by_label.items()
            if any(label.casefold() == alias.casefold() for alias in aliases)
        ]
        if exact:
            label, cats = exact[0]
            selected.append((cats[0], label, display if label.startswith("Metric: ") else label))
    return list(dict.fromkeys(selected))


def _write_summary(
    ws: Any,
    frame: pd.DataFrame,
    periods: list[Any],
    lineage_rows: Mapping[tuple[str, str, str], int],
    lineage_items: Mapping[tuple[str, str, str], Mapping[str, Any]],
    title: str,
    note: str,
) -> None:
    _write_title(ws, title, note)
    if not periods:
        ws["B5"] = "No matching periods are present in this build."
        ws["B5"].font = _BODY
        ws.column_dimensions["B"].width = 72
        return
    _write_header(ws, 5, ["Metric", *[_text(p) for p in periods]])
    row_no = 6
    for category, label, display in _summary_rows(frame):
        ws.cell(row_no, 2, display).font = _BODY
        for column, period in enumerate(periods, 3):
            value = frame.at[(category, label), period]
            if isinstance(value, pd.Series):
                value = value.iloc[0]
            cell = ws.cell(row_no, column)
            key = (category, label, _text(period))
            _write_front_value(cell, value, label, lineage_items.get(key, {}))
            target = lineage_rows.get(key)
            if target and _present(value):
                _set_internal_link(cell, target)
        row_no += 1
    ws.freeze_panes = "C6"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    for column in range(3, len(periods) + 3):
        ws.column_dimensions[get_column_letter(column)].width = 14
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 2
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"
    ws.print_title_cols = "B:B"
    ws.print_area = f"B2:{get_column_letter(len(periods) + 2)}{max(row_no - 1, 6)}"


def _write_table_sheet(ws: Any, title: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
    _write_title(ws, title)
    # Repeated identifier column B must not repeat a clipped fragment of a
    # merged title on horizontal continuation pages. The page header carries
    # the complete title on every page.
    ws['B2'] = None
    ws['C2'] = _safe_text(title)
    ws['C2'].font = _TITLE
    ws.merge_cells('C2:F2')
    ws.oddHeader.center.text = title
    ws.oddHeader.center.size = 12
    _write_header(ws, 4, headers)
    for row_number, row in enumerate(rows, 5):
        for column, header in enumerate(headers, 2):
            cell = ws.cell(row_number, column)
            value = row.get(header, "")
            if header in {"Value", "Raw Reported Value", "Unit Scale", "Display Scale", "Failures"}:
                _write_value(cell, value, header)
            else:
                cell.value = _safe_text(value)
                cell.alignment = _ALIGN_LEFT
            if header in {"Metric", "Finding", "Suggested Action", "Formula", "Context", "XBRL Concept", "Locator", "Method", "Evidence"}:
                cell.alignment = _ALIGN_WRAP_TOP
                if len(_safe_text(value)) > 42:
                    ws.row_dimensions[row_number].height = max(ws.row_dimensions[row_number].height or 15, 30)
            elif cell.alignment.horizontal != "right":
                cell.alignment = _ALIGN_TOP
    end_row = max(5, 4 + len(rows))
    if rows:
        table = Table(displayName=re.sub(r"\W", "", title.title()), ref=f"B4:{get_column_letter(len(headers) + 1)}{end_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"B4:{get_column_letter(len(headers) + 1)}{end_row}"
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 3
    widths = {
        "Metric": 34, "Finding": 52, "Suggested Action": 44, "Formula": 42,
        "Context": 30, "XBRL Concept": 42, "Locator": 42, "SEC Link": 24,
        "Accession": 24, "Method": 34, "Evidence": 52, "Check": 34,
    }
    for index, header in enumerate(headers, 2):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(header, min(max(len(header) + 3, 12), 24))
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 2
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    ws.print_title_cols = "B:B"
    ws.print_area = f"B2:{get_column_letter(len(headers) + 1)}{end_row}"


def _ticker(frame: pd.DataFrame) -> str:
    audit = frame.attrs.get("fact_audit")
    if isinstance(audit, pd.DataFrame) and not audit.empty and "Ticker" in audit.columns:
        values = audit["Ticker"].dropna()
        if not values.empty:
            return _text(values.iloc[0]).upper()
    return ""


def _research_kpis(frame: pd.DataFrame, ticker: str) -> list[tuple[str, str, str]]:
    specs = _TLN_KPI_SPECS if ticker == "TLN" else ()
    available: dict[str, list[str]] = defaultdict(list)
    for category, label in frame.index:
        available[_text(label).casefold()].append(_text(category))
    selected: list[tuple[str, str, str]] = []
    for display, *aliases in specs:
        for alias in aliases:
            categories = available.get(alias.casefold(), [])
            if categories:
                selected.append((categories[0], alias, display))
                break
    return selected


def _display_label_with_unit(display: str, item: Mapping[str, Any]) -> str:
    unit = _text(item.get("Display Unit"))
    readable = {"USD/MWh": "$/MWh", "USD/share": "$/share"}.get(unit, unit)
    return f"{display} ({readable})" if readable else display


def _write_research_kpis(
    ws: Any,
    frame: pd.DataFrame,
    periods: list[Any],
    ticker: str,
    lineage_rows: Mapping[tuple[str, str, str], int],
    lineage_items: Mapping[tuple[str, str, str], Mapping[str, Any]],
) -> None:
    _write_title(
        ws,
        f"{ticker or 'Company'} Operating KPI",
        "Curated research metrics only. Standardized calculations are labeled; raw extraction is in 97_RAW_KPI.",
    )
    _write_header(ws, 5, ["Metric", *[_text(period) for period in periods]])
    row_no = 6
    for category, label, display in _research_kpis(frame, ticker):
        representative = next((
            lineage_items.get((category, label, _text(period)), {})
            for period in reversed(periods)
            if _present(frame.at[(category, label), period])
        ), {})
        ws.cell(row_no, 2, _display_label_with_unit(display, representative)).font = _BODY
        for column, period in enumerate(periods, 3):
            value = frame.at[(category, label), period]
            if isinstance(value, pd.Series):
                value = value.iloc[0]
            key = (category, label, _text(period))
            item = lineage_items.get(key, {})
            cell = ws.cell(row_no, column)
            _write_front_value(cell, value, label, item)
            target = lineage_rows.get(key)
            if target and _present(value):
                _set_internal_link(cell, target)
        row_no += 1
    if row_no == 6:
        ws.cell(row_no, 2, "No company KPI mapping is available for this issuer.").font = _BODY
    ws.freeze_panes = "C6"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 36
    for column in range(3, len(periods) + 3):
        ws.column_dimensions[get_column_letter(column)].width = 13
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"B2:{get_column_letter(len(periods) + 2)}{max(row_no - 1, 6)}"


def _program_qa(
    lineage: list[dict[str, Any]],
    sources: list[dict[str, str]],
    details: list[dict[str, str]],
    active_periods: set[str] | None = None,
) -> list[dict[str, Any]]:
    core = [
        row for row in lineage
        if _text(row.get("_Category")).startswith(("1_", "2_", "3_"))
        and _text(row.get("Metric")) not in _ENGINE_DIAGNOSTIC_LABELS
        and (active_periods is None or _text(row.get("Period")) in active_periods)
    ]
    missing_source = [row for row in core if not _text(row.get("Source ID"))]
    missing_formula = [row for row in core if row.get("Type") in {"C", "D"} and not _text(row.get("Formula"))]
    unsafe_formula = [row for row in missing_formula if _text(row.get("Review Required")) != "YES"]
    missing_unit = [row for row in core if not _text(row.get("Unit"))]
    required_labels = {
        "revenue", "operating income", "net income", "cash & equivalents",
        "total assets", "total liabilities", "total equity",
        "operating cash flow", "investing cash flow", "financing cash flow", "net cash flow",
    }
    blocked_front = [
        row for row in core
        if _text(row.get("Metric")).casefold() in required_labels and not _front_value_is_eligible(row)
    ]

    by_period: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in core:
        by_period[_text(row.get("Period"))][_text(row.get("Metric")).casefold()] = row

    def equation_check(
        aliases: tuple[tuple[str, ...], ...],
        expression: Any,
    ) -> tuple[str, int, str]:
        tested = 0
        failures = 0
        for metrics in by_period.values():
            rows = [next((metrics.get(alias.casefold()) for alias in group if metrics.get(alias.casefold())), None) for group in aliases]
            if any(row is None for row in rows):
                continue
            values = [_decimal(row.get("Value")) for row in rows if row is not None]
            if len(values) != len(rows) or any(value is None for value in values):
                continue
            tested += 1
            resolved = [value for value in values if value is not None]
            difference = abs(expression(resolved))
            tolerance = max(Decimal("1"), max(abs(value) for value in resolved) * Decimal("0.000001"))
            if difference > tolerance:
                failures += 1
        if not tested:
            return "NOT_RUN", 1, "Required line items were not available in a common period."
        return ("PASS" if not failures else "FAIL"), failures, f"Evaluated {tested} period(s) using Decimal arithmetic."

    balance_status, balance_failures, balance_evidence = equation_check(
        (("Total Assets",), ("Total Liabilities",), ("Total Equity", "Stockholders Equity")),
        lambda values: values[0] - values[1] - values[2],
    )
    cash_status, cash_failures, cash_evidence = equation_check(
        (("Operating Cash Flow",), ("Investing Cash Flow",), ("Financing Cash Flow",), ("Net Cash Flow",)),
        lambda values: values[0] + values[1] + values[2] - values[3],
    )
    checks = [
        ("QA001", "Core lineage coverage", "PASS" if not missing_source else "FAIL", len(missing_source), "Every core value has a normalized filing source."),
        ("QA002", "Calculated-value safety", "PASS" if not unsafe_formula else "FAIL", len(unsafe_formula), f"{len(missing_formula)} calculated or derived value(s) without an expression are withheld from front sheets."),
        ("QA003", "Core unit coverage", "PASS" if not missing_unit else "FAIL", len(missing_unit), "Every core value has a recognized unit."),
        ("QA004", "Source normalization", "PASS" if sources else "FAIL", 0 if sources else 1, "Filing metadata is stored once and referenced by Source ID."),
        (
            "QA005",
            "Front-value eligibility",
            "WARNING" if blocked_front else "PASS",
            len(blocked_front),
            f"{len(blocked_front)} unresolved core value(s) remain in Lineage and render as unavailable in front sheets.",
        ),
        ("QA006", "Balance sheet equation", balance_status, balance_failures, balance_evidence),
        ("QA007", "Cash flow subtotal", cash_status, cash_failures, cash_evidence),
        ("QA008", "Presentation separation", "PASS", 0, "Raw KPI and detailed QA rows are isolated in hidden backend sheets."),
    ]
    rows: list[dict[str, Any]] = []
    for check_id, check, status, failures, evidence in checks:
        rows.append({
            "Check ID": check_id,
            "Status": status,
            "Severity": "INFO" if status == "PASS" else "HIGH",
            "Check": check,
            "Failures": failures,
            "Evidence": evidence,
        })
    if details:
        rows.append({
            "Check ID": "QA009",
            "Status": "WARNING",
            "Severity": "MEDIUM",
            "Check": "Raw extraction diagnostics",
            "Failures": len(details),
            "Evidence": "See 98_RAW_QA. These program diagnostics are not sent to AI review.",
        })
    return rows


def _semantic_review_queue(
    lineage: list[dict[str, Any]],
    front_periods: set[str],
    front_metrics: set[tuple[str, str]],
) -> list[dict[str, Any]]:
    """Return a small, deterministic queue of economic/semantic anomalies."""
    findings: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    per_unit_anomalies: dict[str, list[dict[str, Any]]] = defaultdict(list)
    percentage_anomalies: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)

    def add(severity: str, row: Mapping[str, Any], check: str, finding: str) -> None:
        key = (_text(row.get("Metric")), _text(row.get("Period")), check)
        if key in seen:
            return
        seen.add(key)
        value = _number(row.get("Value"))
        display_scale = float(_number(row.get("Display Scale")) or 1.0)
        findings.append({
            "Status": "OPEN",
            "Severity": severity,
            "Metric": key[0],
            "Period": key[1],
            "Value": float(value) * display_scale if value is not None else None,
            "Unit": row.get("Display Unit") or row.get("Unit"),
            "Source ID": row.get("Source ID"),
            "Check": check,
            "Finding": finding,
            "Suggested Action": "Review the linked lineage and filing, then record a concise conclusion.",
        })

    quarterly: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in lineage:
        metric_key = (_text(row.get("_Category")), _text(row.get("Metric")))
        if _text(row.get("Period")) not in front_periods or metric_key not in front_metrics:
            continue
        value = _number(row.get("Value"))
        if value is None:
            continue
        unit = _text(row.get("Unit"))
        metric = _text(row.get("Metric"))
        period = _text(row.get("Period"))
        display_scale = float(_number(row.get("Display Scale")) or 1.0)
        normalized_value = float(value) * display_scale
        method = _text(row.get("Method"))
        deterministic_withholds = {
            "MISSING_RELOCATABLE_REPORTED_FACT", "REJECTED_REVENUE_AS_EXPENSE",
            "REJECTED_OPERATING_INCOME_AS_EXPENSE", "UNPROVEN_DEBT_FAMILY_SCOPE",
            "REJECTED_NON_ADDITIVE_PERIOD_MISMATCH",
        }
        known_missing_derivation = (
            method == "MISSING_RECURSIVE_OPERANDS" and (
                metric in {"Net Long-Term Debt Issued (Repaid)", "Shares Issued - Common Stock", "Shares Issued - Preferred Stock"}
                or metric.startswith("Other Costs and Expenses - ")
            )
        ) or metric in {
            "Net Shares Issued (Repurchased)",
            "Other Financing Adjustments (Net)", "Other Investing Adjustments (Net)",
        }
        if _text(row.get("Review Required")) == "YES":
            if method not in deterministic_withholds and not known_missing_derivation:
                add(
                    "HIGH",
                    row,
                    "Low-confidence classification",
                    "Value withheld from financial front; evidence rule: " + method,
                )
            continue  # Rejected candidates cannot also become numeric anomaly signals.
        if _text(row.get("Type")) == "D" and not _text(row.get("Formula")):
            add(
                "HIGH",
                row,
                "Missing derivation formula",
                "A visible derived value has no retained derivation expression.",
            )
        if unit == "USD/MWh" and abs(normalized_value) > 10_000:
            per_unit_anomalies[metric].append(row)
        if unit == "%" and abs(float(value)) > 100:
            percentage_anomalies[(metric, "Percentage outside expected range")].append(row)
        elif unit == "%" and "gross margin" in metric.casefold() and abs(float(value)) >= 99.9:
            percentage_anomalies[(metric, "Economic meaning review")].append(row)
        if re.fullmatch(r"\d{4}-Q[1-4]", period):
            quarterly[metric].append(row)

    for metric, rows in per_unit_anomalies.items():
        current = max(rows, key=lambda row: _period_key(row.get("Period")))
        value = _number(current.get("Value"))
        raw = _text(current.get("Raw Reported Value"))
        add(
            "CRITICAL",
            current,
            "Possible unit scaling anomaly",
            f"{len(rows)} period(s) are affected; latest stored value {float(value or 0):g} USD/MWh versus source display {raw or 'not recorded'}.",
        )

    for (metric, check), rows in percentage_anomalies.items():
        current = max(rows, key=lambda row: _period_key(row.get("Period")))
        value = _number(current.get("Value"))
        finding = (
            f"{len(rows)} period(s) exceed the expected percentage range; latest is {float(value or 0):g}%."
            if check == "Percentage outside expected range"
            else f"{len(rows)} period(s) show a 100% gross margin; review whether costs are classified elsewhere."
        )
        add("HIGH" if check.startswith("Percentage") else "MEDIUM", current, check, finding)

    for metric, rows in quarterly.items():
        ordered = sorted(rows, key=lambda row: _period_key(row.get("Period")))
        outliers: list[tuple[float, dict[str, Any], dict[str, Any]]] = []
        for previous, current in zip(ordered, ordered[1:]):
            previous_ordinal = _quarter_ordinal(previous.get("Period"))
            current_ordinal = _quarter_ordinal(current.get("Period"))
            if previous_ordinal is None or current_ordinal != previous_ordinal + 1:
                continue
            a, b = _number(previous.get("Value")), _number(current.get("Value"))
            if a is None or b is None or abs(float(a)) < 1:
                continue
            change = abs((float(b) - float(a)) / float(a))
            if change >= 4 and max(abs(float(a)), abs(float(b))) >= 50_000_000:
                outliers.append((change, previous, current))
        if outliers:
            change, previous, current = max(outliers, key=lambda item: item[0])
            add(
                "MEDIUM",
                current,
                "Quarter-over-quarter outlier",
                f"{len(outliers)} flagged transition(s); largest changed {change:.1f}x versus {_text(previous.get('Period'))}.",
            )

    severity_rank = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "INFO": 3}
    findings.sort(key=lambda row: (severity_rank.get(_text(row.get("Severity")), 9), _text(row.get("Metric")), _text(row.get("Period"))))
    for number, row in enumerate(findings, 1):
        row["Review ID"] = f"R{number:03d}"
    return findings


def _write_overview(
    ws: Any,
    frame: pd.DataFrame,
    periods: list[Any],
    lineage: list[dict[str, Any]],
    sources: list[dict[str, str]],
    review_items: list[dict[str, Any]],
    qa_rows: list[dict[str, Any]],
    lineage_rows: Mapping[tuple[str, str, str], int],
    lineage_items: Mapping[tuple[str, str, str], Mapping[str, Any]],
) -> None:
    ticker = _ticker(frame)
    _write_title(ws, f"{ticker or 'SEC'} Financial Workbook", "Clean financial views backed by normalized SEC sources and cell-level lineage.")
    ws.merge_cells('B2:G2')
    required_core = {"QA001", "QA002", "QA003", "QA004", "QA005", "QA006", "QA007"}
    core_results = {
        _text(row.get("Check ID")): _text(row.get("Status")) for row in qa_rows
    }
    if not required_core.issubset(core_results):
        core_status = "REVIEW"
    elif all(core_results[key] == "PASS" for key in required_core):
        core_status = "PASS"
    elif any(core_results[key] == "FAIL" for key in required_core):
        core_status = "FAIL"
    else:
        core_status = "REVIEW"
    company_kpi_status = (
        "CONFIGURED" if ticker in _INDUSTRY_COVERAGE else "NOT CONFIGURED"
    )
    labels = [
        ("Financial Core", core_status),
        ("Industry Mapping", _INDUSTRY_COVERAGE.get(ticker, "PARTIAL")),
        ("Company KPI", company_kpi_status),
        ("Company", ticker or "See source filings"),
        ("Periods", f"{len(periods)} ({_text(periods[0]) if periods else 'none'} to {_text(periods[-1]) if periods else 'none'})"),
        ("Sources", len(sources)),
        ("Lineage records", len(lineage)),
        ("Open semantic review items", len(review_items)),
    ]
    for row, (label, value) in enumerate(labels, 5):
        ws.cell(row, 2, label).font = _BOLD
        ws.cell(row, 3, value).font = _BODY
    ws["B15"] = "Navigation"
    ws["B15"].font = _BOLD
    for row, sheet in enumerate(SHEET_NAMES[1:10], 16):
        cell = ws.cell(row, 2, sheet)
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1")
        cell.font = _LINK
    ws.cell(26, 2, "Audit backend: unhide 90_SOURCES, 91_LINEAGE, 97_RAW_KPI or 98_RAW_QA when needed.")
    ws.cell(26, 2).font = Font(name="Arial", size=9, italic=True, color="666666")
    if periods:
        latest = periods[-1]
        ws["E5"] = f"Latest period: {_text(latest)}"
        ws["E5"].font = _BOLD
        ws["E6"] = "GAAP / standardized"
        ws["E6"].fill = _SECTION
        ws["E6"].font = _BOLD
        output_row = 7
        for category, label, display in _summary_rows(frame):
            value = frame.at[(category, label), latest]
            if isinstance(value, pd.Series):
                value = value.iloc[0]
            if not _present(value):
                continue
            key = (category, label, _text(latest))
            item = lineage_items.get(key, {})
            ws.cell(output_row, 5, _display_label_with_unit(display, item)).font = _BODY
            cell = ws.cell(output_row, 6)
            _write_front_value(cell, value, label, item)
            target = lineage_rows.get(key)
            if target:
                _set_internal_link(cell, target)
            output_row += 1
        output_row += 1
        ws.cell(output_row, 5, "Company / operating KPIs").fill = _SECTION
        ws.cell(output_row, 5).font = _BOLD
        output_row += 1
        for category, label, display in _research_kpis(frame, ticker)[:5]:
            value = frame.at[(category, label), latest]
            if isinstance(value, pd.Series):
                value = value.iloc[0]
            if not _present(value):
                continue
            key = (category, label, _text(latest))
            item = lineage_items.get(key, {})
            ws.cell(output_row, 5, _display_label_with_unit(display, item)).font = _BODY
            cell = ws.cell(output_row, 6)
            _write_front_value(cell, value, label, item)
            target = lineage_rows.get(key)
            if target:
                _set_internal_link(cell, target)
            output_row += 1
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["E"].width = 34
    ws.column_dimensions["F"].width = 18
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"B2:F{max(28, output_row if periods else 28)}"


def _combine_frames(frames: Iterable[pd.DataFrame]) -> pd.DataFrame:
    usable = [frame for frame in frames if isinstance(frame, pd.DataFrame) and not frame.empty]
    if not usable:
        return pd.DataFrame()
    matrices = [pd.DataFrame(frame.to_numpy(copy=False), index=frame.index, columns=frame.columns)
                for frame in usable]
    combined = matrices[0].copy()
    for frame in matrices[1:]:
        combined = combined.combine_first(frame)
        for column in frame.columns:
            if column not in combined.columns:
                combined[column] = frame[column]
    combined.attrs.update(usable[0].attrs)
    audits = [frame.attrs.get("fact_audit") for frame in usable]
    audits = [audit for audit in audits if isinstance(audit, pd.DataFrame) and not audit.empty]
    combined.attrs["fact_audit"] = pd.concat(audits, ignore_index=True, sort=False) if audits else pd.DataFrame()
    dates: dict[Any, Any] = {}
    for frame in usable:
        dates.update(frame.attrs.get("period_dates") or {})
    combined.attrs["period_dates"] = dates
    return combined


def save_pivot_xlsx(
    final_pivot: pd.DataFrame,
    out_path: str | Path,
    annual_pivot: pd.DataFrame | None = None,
) -> Path:
    """Write clean research views plus normalized SEC source and lineage ledgers."""
    _canonical_categories(final_pivot)
    primary = final_pivot.copy()
    primary.attrs = dict(final_pivot.attrs)
    primary_periods, primary_dates = _periods(primary)
    primary = primary.loc[:, primary_periods]

    primary_is_annual = bool(primary_periods) and all(
        re.search(r"(?:^FY|FY$)", _text(period), re.I) for period in primary_periods
    )
    annual = annual_pivot.copy() if isinstance(annual_pivot, pd.DataFrame) and not annual_pivot.empty else None
    if annual is not None:
        annual.attrs = dict(annual.attrs)
    elif primary_is_annual:
        annual = primary
    annual_periods: list[Any] = []
    if annual is not None:
        annual_periods, _ = _periods(annual)
        annual_periods = [period for period in annual_periods if re.search(r"(?:^FY|FY$)", _text(period), re.I)][-5:]
        annual = annual.loc[:, annual_periods]

    quarter_periods = [period for period in primary_periods if re.search(r"Q[1-4]$", _text(period), re.I)][-12:]
    detail_periods = quarter_periods if quarter_periods else primary_periods
    combined = _combine_frames([primary, annual] if annual is not None else [primary])
    combined_periods, _ = _periods(combined)
    audit = combined.attrs.get("fact_audit")
    if not isinstance(audit, pd.DataFrame):
        audit = pd.DataFrame()
    sources, accession_to_id = _source_rows(audit)
    lineage, qa_details = _build_lineage(combined, audit, combined_periods, accession_to_id)
    render_links = {
        (_text(item["_Category"]), _text(item["Metric"]), _text(item["Period"])): index
        for index, item in enumerate(lineage, 5)
    }
    lineage_items = {
        (_text(item["_Category"]), _text(item["Metric"]), _text(item["Period"])): item
        for item in lineage
    }
    ticker = _ticker(primary) or (_ticker(annual) if annual is not None else "")
    front_metrics = {
        (category, label) for category, label, _display in _summary_rows(primary)
    }
    front_metrics.update(
        (category, label) for category, label, _display in _research_kpis(primary, ticker)
    )
    visible_categories = {
        category
        for categories in _MAIN_CATEGORIES.values()
        for category in categories
    }
    front_metrics.update(
        (_text(category), _text(label))
        for category, label in primary.index
        if _text(category) in visible_categories
    )
    review_items = _semantic_review_queue(
        lineage,
        {_text(period) for period in detail_periods},
        front_metrics,
    )
    qa_rows = _program_qa(
        lineage,
        sources,
        qa_details,
        active_periods={_text(period) for period in [*detail_periods, *annual_periods]},
    )

    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEET_NAMES:
        wb.create_sheet(name)

    _write_overview(
        wb["00_OVERVIEW"], primary, detail_periods, lineage, sources,
        review_items, qa_rows, render_links, lineage_items,
    )
    if annual is not None:
        _write_summary(
            wb["01_ANNUAL"], annual, annual_periods, render_links, lineage_items,
            "Annual Summary", "USD/shares: millions; EPS: original units.",
        )
    else:
        _write_title(wb["01_ANNUAL"], "Annual Summary")
        wb["01_ANNUAL"]["B5"] = "Annual companion data was not produced by this build."
    _write_summary(
        wb["02_QUARTERLY"], primary, quarter_periods, render_links, lineage_items,
        "Quarterly Summary", "USD/shares: millions; EPS: original units.",
    )
    for sheet, categories in _MAIN_CATEGORIES.items():
        _write_matrix(
            wb[sheet], primary, detail_periods, primary_dates, categories,
            render_links, lineage_items, sheet[3:].replace("_", " ").title(),
        )
    _write_research_kpis(
        wb["07_OPERATING_KPI"], primary, quarter_periods, ticker,
        render_links, lineage_items,
    )

    qa_headers = ["Check ID", "Status", "Severity", "Check", "Failures", "Evidence"]
    _write_table_sheet(wb["08_QA"], "Program QA", qa_rows, qa_headers)
    review_headers = [
        "Review ID", "Status", "Severity", "Metric", "Period", "Value", "Unit",
        "Source ID", "Check", "Finding", "Suggested Action",
    ]
    _write_table_sheet(wb["09_REVIEW"], "Semantic Review Queue", review_items, review_headers)

    source_headers = [
        "Source ID", "Ticker", "Form", "Report Period", "Filed", "Accepted At",
        "Accession", "SEC Link", "Evidence Class",
    ]
    _write_table_sheet(wb["90_SOURCES"], "SEC Sources", sources, source_headers)
    for row_number, row in enumerate(sources, 5):
        url = row.get("SEC Link", "")
        if url:
            cell = wb["90_SOURCES"].cell(row_number, source_headers.index("SEC Link") + 2)
            cell.value = "SEC 原文 ↗"
            cell.hyperlink = url
            cell.font = _LINK

    lineage_headers = list(_LINEAGE_EXPORT_FIELDS)
    _write_table_sheet(wb["91_LINEAGE"], "Cell Lineage", lineage, lineage_headers)
    source_lookup = {row["Source ID"]: index + 5 for index, row in enumerate(sources)}
    for row_number, item in enumerate(lineage, 5):
        source_row = source_lookup.get(_text(item.get("Source ID")))
        if source_row:
            cell = wb["91_LINEAGE"].cell(row_number, lineage_headers.index("Source ID") + 2)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'90_SOURCES'!B{source_row}")
            cell.font = _LINK

    raw_kpi = [
        {header: item.get(header, "") for header in lineage_headers}
        for item in lineage if _text(item.get("_Category")) in _RAW_KPI_CATEGORIES
    ]
    _write_table_sheet(wb["97_RAW_KPI"], "Raw KPI Extraction", raw_kpi, lineage_headers)

    raw_qa: list[dict[str, Any]] = []
    for number, issue in enumerate(qa_details, 1):
        raw_qa.append({"Record ID": f"Q{number:05d}", **issue})
    for item in lineage:
        if _text(item.get("_Category")) in _RAW_QA_CATEGORIES:
            raw_qa.append({
                "Record ID": f"Q{len(raw_qa) + 1:05d}", "Status": "SOURCE",
                "Severity": "INFO", "Metric": item.get("Metric"), "Period": item.get("Period"),
                "Source ID": item.get("Source ID"), "Check": "Engine integrity output",
                "Finding": item.get("Value"), "Suggested Action": "Use as program diagnostic evidence.",
            })
    raw_qa_headers = [
        "Record ID", "Status", "Severity", "Metric", "Period", "Source ID",
        "Check", "Finding", "Suggested Action",
    ]
    _write_table_sheet(wb["98_RAW_QA"], "Raw Program QA", raw_qa, raw_qa_headers)

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = "17365D" if not ws.title.startswith("9") else "7F8C8D"
        ws.sheet_view.zoomScale = 90
        ws.oddHeader.center.text = f"&B{ws.title}"
        ws.oddFooter.right.text = "Page &P of &N"
    for name in ("90_SOURCES", "91_LINEAGE", "97_RAW_KPI", "98_RAW_QA"):
        wb[name].sheet_state = "hidden"

    destination = Path(out_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    _write_build_manifest(destination, ticker, sources, lineage, qa_rows)
    return destination


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical_scalar(value: Any, *, numeric: bool = False) -> Any:
    if value is None or not _present(value):
        return None
    if isinstance(value, bool):
        return value
    parsed = _decimal(value)
    if parsed is not None and (numeric or not isinstance(value, str)):
        return format(parsed.normalize(), "f")
    return _text(value)


def _semantic_projection(
    ticker: str,
    sources: list[dict[str, str]],
    lineage: list[dict[str, Any]],
    qa_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    def rows(items: list[dict[str, Any]], fields: tuple[str, ...]) -> list[dict[str, Any]]:
        numeric_fields = {"Value", "Raw Reported Value", "Unit Scale", "Display Scale", "Failures"}
        return [
            {
                field: _canonical_scalar(
                    item.get(field),
                    numeric=field in numeric_fields,
                )
                for field in fields
            }
            for item in items
        ]

    return {
        "schema": "sec-data-build-v1",
        "ticker": ticker,
        "sources": rows(sorted(sources, key=lambda item: _text(item.get("Source ID"))), _SOURCE_EXPORT_FIELDS),
        "lineage": rows(sorted(lineage, key=lambda item: _text(item.get("Lineage ID"))), _LINEAGE_EXPORT_FIELDS),
        "qa": rows(sorted(qa_rows, key=lambda item: _text(item.get("Check ID"))), _QA_EXPORT_FIELDS),
    }


def _workbook_audit_rows(
    workbook_path: Path,
    book: Any | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    book = book or load_workbook(workbook_path, data_only=True, read_only=False)

    def read_rows(sheet: str, fields: tuple[str, ...], key: str) -> list[dict[str, Any]]:
        ws = book[sheet]
        header_map = {
            _text(ws.cell(4, column).value): column
            for column in range(2, ws.max_column + 1)
        }
        if not set(fields).issubset(header_map):
            raise ValueError(f"{sheet} headers are incomplete")
        output: list[dict[str, Any]] = []
        for row_number in range(5, ws.max_row + 1):
            if not _text(ws.cell(row_number, header_map[key]).value):
                continue
            item: dict[str, Any] = {}
            for field in fields:
                cell = ws.cell(row_number, header_map[field])
                value = cell.value
                if field == "SEC Link" and cell.hyperlink is not None:
                    value = cell.hyperlink.target
                if (
                    isinstance(value, str)
                    and value.startswith("'")
                    and value[1:].startswith(("=", "+", "-", "@"))
                ):
                    value = value[1:]
                item[field] = value
            output.append(item)
        return output

    return (
        read_rows("90_SOURCES", _SOURCE_EXPORT_FIELDS, "Source ID"),
        read_rows("91_LINEAGE", _LINEAGE_EXPORT_FIELDS, "Lineage ID"),
        read_rows("08_QA", _QA_EXPORT_FIELDS, "Check ID"),
    )


def _write_build_manifest(
    workbook_path: Path,
    ticker: str,
    sources: list[dict[str, str]],
    lineage: list[dict[str, Any]],
    qa_rows: list[dict[str, Any]],
) -> Path:
    workbook_sources, workbook_lineage, workbook_qa = _workbook_audit_rows(workbook_path)
    projection = _semantic_projection(
        ticker,
        workbook_sources,
        workbook_lineage,
        workbook_qa,
    )
    semantic_bytes = json.dumps(
        projection,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    code_files = [
        Path(__file__),
        Path(__file__).with_name("sec_data.py"),
        Path(__file__).with_name("sec_data_cli.py"),
        Path(__file__).with_name("pipeline_audit.py"),
    ]
    code_digest = hashlib.sha256()
    for path in sorted((path for path in code_files if path.exists()), key=lambda item: item.name):
        code_digest.update(path.name.encode("utf-8"))
        code_digest.update(path.read_bytes())
    required_core = {"QA001", "QA002", "QA003", "QA004", "QA005", "QA006", "QA007"}
    core_results = {
        _text(row.get("Check ID")): _text(row.get("Status"))
        for row in qa_rows
    }
    if not required_core.issubset(core_results):
        core_status = "REVIEW"
    elif all(core_results[check_id] == "PASS" for check_id in required_core):
        core_status = "PASS"
    elif any(core_results[check_id] == "FAIL" for check_id in required_core):
        core_status = "FAIL"
    else:
        core_status = "REVIEW"
    manifest = {
        "schema": "sec-data-build-manifest-v1",
        "artifact_status": "DRAFT",
        "financial_core_status": core_status,
        "ticker": ticker,
        "python": platform.python_version(),
        "source_code_sha256": code_digest.hexdigest(),
        "semantic_sha256": hashlib.sha256(semantic_bytes).hexdigest(),
        "front_bindings": _front_binding_contract(workbook_path),
        "workbook": {
            "file": workbook_path.name,
            "sha256": _sha256_file(workbook_path),
            "bytes": workbook_path.stat().st_size,
        },
        "source_accessions": [row.get("Accession", "") for row in workbook_sources],
        "counts": {
            "sources": len(workbook_sources),
            "lineage": len(workbook_lineage),
            "qa": len(workbook_qa),
        },
    }
    manifest_path = workbook_path.with_suffix(".manifest.json")
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest_path


def verify_workbook_manifest(workbook_path: str | Path) -> list[str]:
    """Return deterministic integrity errors for an exported workbook."""
    workbook = Path(workbook_path)
    manifest_path = workbook.with_suffix(".manifest.json")
    if not workbook.is_file():
        return ["workbook is missing"]
    if not manifest_path.is_file():
        return ["build manifest is missing"]
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return ["build manifest is invalid"]
    if manifest.get("schema") != "sec-data-build-manifest-v1":
        return ["build manifest schema is unsupported"]
    expected = _text((manifest.get("workbook") or {}).get("sha256"))
    if not expected or expected != _sha256_file(workbook):
        return ["workbook hash does not match build manifest"]
    book = None
    try:
        book = load_workbook(workbook, data_only=True, read_only=False)
        sources, lineage, qa_rows = _workbook_audit_rows(workbook, book)
        projection = _semantic_projection(_text(manifest.get("ticker")), sources, lineage, qa_rows)
        semantic_bytes = json.dumps(
            projection,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        if hashlib.sha256(semantic_bytes).hexdigest() != _text(manifest.get("semantic_sha256")):
            return ["workbook semantic content does not match build manifest"]
        expected_counts = manifest.get("counts") or {}
        actual_counts = {"sources": len(sources), "lineage": len(lineage), "qa": len(qa_rows)}
        if any(expected_counts.get(key) != value for key, value in actual_counts.items()):
            return ["workbook row counts do not match build manifest"]
        front_errors = _verify_front_values(workbook, book)
        if front_errors:
            return front_errors
        if manifest.get('front_bindings') != _front_binding_contract(workbook, book):
            return ['front values, labels, periods or lineage bindings differ from the frozen build']
    except (KeyError, OSError, ValueError):
        return ["workbook audit sheets are invalid"]
    finally:
        if book is not None:
            book.close()
    return []


def _front_binding_contract(path: Path, book: Any | None = None) -> list[dict[str, Any]]:
    owned = book is None
    book = book or load_workbook(path, data_only=True)
    try:
        result = []
        for name in SHEET_NAMES[:8]:
            for row in book[name].iter_rows():
                for cell in row:
                    link = _text(cell.hyperlink.location) if cell.hyperlink else ''
                    if '91_LINEAGE' in link:
                        result.append({'sheet': name, 'cell': cell.coordinate, 'binding': link,
                                       'value': _canonical_scalar(cell.value, numeric=True),
                                       'label': _text(book[name].cell(cell.row, 5 if name == '00_OVERVIEW' else 2).value),
                                       'period': _text(book[name].cell(5, cell.column).value) if name != '00_OVERVIEW' else ''})
        return result
    finally:
        if owned:
            book.close()


def _verify_front_values(path: Path, book: Any | None = None) -> list[str]:
    """Reconcile displayed amounts to the exact background row they link to."""
    owned = book is None
    book = book or load_workbook(path, data_only=True)
    try:
        audit = book["91_LINEAGE"]
        last_audit_row = audit.max_row
        headers = {_text(audit.cell(4, col).value): col for col in range(2, audit.max_column + 1)}
        for name in SHEET_NAMES[:8]:
            for row in book[name].iter_rows(min_row=5, min_col=3):
                for cell in row:
                    link = _text(cell.hyperlink.location) if cell.hyperlink else ""
                    match = re.fullmatch(r"'91_LINEAGE'!A(\d+)", link)
                    if not match:
                        if isinstance(cell.value, (int, float)) and not (
                            name == '00_OVERVIEW' and cell.coordinate in {'C10', 'C11', 'C12'}
                        ):
                            return [f"{name}!{cell.coordinate} has no lineage binding"]
                        continue
                    target = int(match[1])
                    if target < 5 or target > last_audit_row:
                        return [f"{name}!{cell.coordinate} has an invalid lineage binding"]
                    item = {field: audit.cell(target, col).value for field, col in headers.items()}
                    if _front_value_is_eligible(item):
                        expected = float(_number(item.get('Value')) or 0) * float(
                            _number(item.get('Display Scale')) or 1)
                        valid = _same_number(cell.value, expected)
                    else:
                        valid = cell.value == '—'
                    if not valid:
                        return [f"{name}!{cell.coordinate} differs from its lineage value"]
        return []
    finally:
        if owned:
            book.close()
