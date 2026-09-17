"""Read-only independent review for a generated SEC financial workbook.

The script never edits the source workbook. It parses cached SEC extracted
instances independently, reconciles visible cells to their lineage rows, and
writes only review artifacts under the selected output directory.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, localcontext
from functools import lru_cache
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
OUT = ROOT / "output" / "reviews"
MAIN_SHEETS = [
    "00_OVERVIEW", "01_ANNUAL", "02_QUARTERLY", "03_INCOME_STATEMENT",
    "04_BALANCE_SHEET", "05_CASH_FLOW", "06_SEGMENTS", "07_OPERATING_KPI",
]
FINANCIAL_SHEETS = set(MAIN_SHEETS[1:])
SOURCE_HEADERS = ("Source ID", "Ticker", "Form", "Report Period", "Filed", "Accepted At", "Accession", "SEC Link", "Evidence Class")
LINEAGE_HEADERS = (
    "Lineage ID", "Excel Sheet", "Metric", "Period", "Value", "Raw Reported Value", "Type",
    "Source ID", "Input Source IDs", "XBRL Concept", "Context", "Unit", "Unit Scale",
    "Display Scale", "Display Unit", "Formula", "Input Facts", "Method", "Locator",
    "Semantic Classification", "Semantic Confidence", "Review Required", "Table Evidence",
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def dec(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value).replace(",", "").replace("(", "-").replace(")", ""))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def close(a: Any, b: Any) -> bool:
    da, db = dec(a), dec(b)
    if da is None or db is None:
        return False
    # XLSX numeric cells are IEEE-754 doubles.  Allow only serialization noise;
    # this tolerance is far below one dollar after a USD-million transform.
    tolerance = max(Decimal("1e-9"), abs(db) * Decimal("1e-12"))
    return abs(da - db) <= tolerance


def parse_period(spec: str) -> tuple[str, str, str] | None:
    spec = spec.split(" | ")[0]
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", spec):
        return ("instant", spec, "")
    spec = spec.split(" | ")[0]
    m = re.fullmatch(r"(\d{4}-\d{2}-\d{2})\.\.(\d{4}-\d{2}-\d{2})", spec)
    return ("duration", m.group(1), m.group(2)) if m else None


def source_rows(book: Any) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    ws = book["90_SOURCES"]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=5, values_only=False):
        row = {header: values[index + 1].value for index, header in enumerate(SOURCE_HEADERS)}
        link = values[8].hyperlink.target if values[8].hyperlink else None
        row["SEC Link"] = link or row["SEC Link"]
        if row["Source ID"]:
            rows.append(row)
    return rows, {text(row["Source ID"]): row for row in rows}


def lineage_rows(book: Any) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    ws = book["91_LINEAGE"]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=5, values_only=True):
        row = {header: values[index + 1] for index, header in enumerate(LINEAGE_HEADERS)}
        if row["Lineage ID"]:
            rows.append(row)
    return rows, {text(row["Lineage ID"]): row for row in rows}


@lru_cache(maxsize=128)
def xml_cache_for(accession: str) -> Path | None:
    data_root = Path(
        os.environ.get("EDGAR_LOCAL_DATA_DIR", str(Path.home() / ".edgar"))
    )
    cache = data_root / "_tcache" / "www.sec.gov"
    compact = accession.replace("-", "")
    candidates = [
        p for p in cache.rglob(f"*{compact}*htm-*")
        if not p.name.endswith(".meta")
    ]
    return candidates[0] if len(candidates) == 1 else None


@lru_cache(maxsize=32)
def parsed_document(path: Path, mtime_ns: int, size: int) -> tuple[bytes, Any]:
    del mtime_ns, size  # Cache identity; each review invocation has a fresh cache.
    raw = path.read_bytes()
    try:
        raw = gzip.decompress(raw)
    except OSError:
        pass
    root = etree.fromstring(raw, parser=etree.XMLParser(resolve_entities=False, no_network=True))
    return raw, root


def xml_facts(path: Path) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    stat = path.stat()
    raw, root = parsed_document(path, stat.st_mtime_ns, stat.st_size)
    contexts: dict[str, dict[str, Any]] = {}
    for node in root.xpath("//*[local-name()='context']"):
        identifier = text(node.xpath("string(.//*[local-name()='identifier'])"))
        instant = text(node.xpath("string(.//*[local-name()='instant'])"))
        start = text(node.xpath("string(.//*[local-name()='startDate'])"))
        end = text(node.xpath("string(.//*[local-name()='endDate'])"))
        dims = [text(x.text) for x in node.xpath(".//*[local-name()='explicitMember']")]
        dims += [text(x.text) for x in node.xpath(".//*[local-name()='typedMember']")]
        contexts[text(node.get("id"))] = {
            "entity": identifier, "kind": "instant" if instant else "duration",
            "start": instant or start, "end": "" if instant else end, "dimensions": dims,
            "axes": [text(x.get("dimension")) for x in node.xpath(".//*[local-name()='explicitMember' or local-name()='typedMember']")],
        }
    units: dict[str, str] = {}
    for node in root.xpath("//*[local-name()='unit']"):
        measures = [text(x.text).split(":")[-1] for x in node.xpath(".//*[local-name()='measure']")]
        units[text(node.get("id"))] = "/".join(measures)
    face_tables = set()
    for table in root.xpath("//*[local-name()='table']"):
        names = {text(x.get("name")).split(":")[-1] for x in table.xpath(".//*[@name]")}
        if "OperatingIncomeLoss" in names and any(name.startswith("EarningsPerShare") for name in names):
            face_tables.add(table)
    facts: list[dict[str, Any]] = []
    # Inline XBRL facts live deep inside the HTML table tree.  Restricting
    # this to root children silently returned zero facts for every filing.
    # Parse every numeric ix/non-ix fact carrying both contextRef and unitRef.
    for node in root.xpath("//*[@contextRef and @unitRef]"):
        context_ref = node.get("contextRef")
        if not context_ref or node.get("unitRef") is None or node.text is None:
            continue
        raw_lexical = text(node.text)
        value = dec(raw_lexical)
        if value is None:
            continue
        scale_raw = text(node.get("scale"))
        if scale_raw:
            try:
                with localcontext() as ctx:
                    ctx.prec = max(50, len(value.as_tuple().digits) + abs(int(scale_raw)) + 5)
                    value = value.scaleb(int(scale_raw))
            except (ValueError, InvalidOperation):
                continue
        if text(node.get("sign")) == "-":
            value = -value
        context = contexts.get(context_ref)
        if context is None:
            continue
        reported_name = text(node.get("name"))
        if ":" in reported_name:
            prefix, local_name = reported_name.split(":", 1)
            namespace = node.nsmap.get(prefix, "")
            concept = local_name
        else:
            qname = etree.QName(node)
            namespace = qname.namespace or ""
            concept = qname.localname
        facts.append({
            "concept": concept,
            "namespace": namespace,
            "context": context_ref,
            "entity": context["entity"],
            "kind": context["kind"], "start": context["start"], "end": context["end"],
            "dimensions": context["dimensions"], "unit": units.get(text(node.get("unitRef")), text(node.get("unitRef"))),
            "value": value, "raw_lexical": raw_lexical,
            "scale": scale_raw, "sign": text(node.get("sign")),
            "id": text(node.get("id")), "decimals": text(node.get("decimals")),
            "axes": context["axes"], "xpath": root.getroottree().getpath(node),
            "artifact_sha256": hashlib.sha256(raw).hexdigest(),
            "face_income": any(table in face_tables for table in node.xpath("ancestor::*[local-name()='table'][1]")),
        })
    return contexts, facts


def unit_matches(expected: str, actual: str) -> bool:
    aliases = {"ratio": "pure", "usd/shares": "usd/share", "usd / shares": "usd/share"}
    return aliases.get(expected.casefold(), expected.casefold()) == aliases.get(actual.casefold(), actual.casefold())


def aspect_name(value: str) -> str:
    value = value.split(":")[-1].removeprefix("dim_us-gaap_").removeprefix("dim_srt_")
    value = re.sub(r'^dim_[^_]+_', '', value)
    return re.sub(r"[^a-z0-9]", "", value.casefold()).removesuffix("member")


def source_entity_cik(source: dict[str, Any]) -> str:
    """Return the issuer CIK, which can differ from the accession prefix.

    Some filings are submitted through another SEC registrant/filing-agent CIK
    (Lumentum is a real example).  The EDGAR archive URL still identifies the
    filing entity in ``/Archives/edgar/data/<issuer-cik>/``.  Treating the
    accession prefix as the entity incorrectly rejects every otherwise exact
    fact in those filings.
    """
    link = text(source.get("SEC Link"))
    match = re.search(r"/Archives/edgar/data/(\d+)/", link, re.I)
    if match:
        return match.group(1).lstrip("0")
    return text(source.get("Accession")).split("-")[0].lstrip("0")


def matching_facts(lineage: dict[str, Any], facts_by_source: dict[str, list[dict[str, Any]]], source_map: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    sid = text(lineage.get("Source ID"))
    concept = text(lineage.get("XBRL Concept"))
    period = parse_period(text(lineage.get("Context")))
    if period is None:
        return []
    namespace = None
    if concept.startswith("{") and "}" in concept:
        namespace, concept = concept[1:].split("}", 1)
    elif ":" in concept:
        prefix, concept = concept.split(":", 1)
        if prefix == "us-gaap":
            namespace = "us-gaap"
    dimensions = text(lineage.get("Context")).split(" | ", 1)
    expected_axes: set[str] = set()
    expected_members: set[str] = set()
    if len(dimensions) == 2:
        parts = dimensions[1].split("; ", 1)
        expected_axes = {aspect_name(x) for x in parts[0].split("|") if x}
        expected_members = {aspect_name(x) for x in parts[1].split("|") if x} if len(parts) == 2 else set()
    source = source_map.get(sid, {})
    cik = source_entity_cik(source)
    candidates = []
    for fact in facts_by_source.get(sid, []):
        if fact["concept"] != concept or fact["entity"].lstrip("0") != cik:
            continue
        if namespace and namespace != fact["namespace"] and not (namespace == "us-gaap" and "/us-gaap/" in fact["namespace"]):
            continue
        if (fact["kind"], fact["start"], fact["end"]) != period or not unit_matches(text(lineage.get("Unit")), fact["unit"]):
            continue
        if {aspect_name(x) for x in fact.get("axes", [])} != expected_axes or {aspect_name(x) for x in fact["dimensions"]} != expected_members:
            continue
        candidates.append(fact)
    if text(lineage.get("Metric")) == "Revenue":
        candidates = [fact for fact in candidates if fact.get("face_income")]
    return candidates


def rounding_interval(fact: dict[str, Any]) -> tuple[Decimal, Decimal]:
    value = fact["value"]
    try:
        radius = Decimal(5).scaleb(-int(fact.get("decimals", "INF")) - 1)
    except ValueError:
        radius = Decimal(0)
    return value - radius, value + radius


def xbrl_match(lineage: dict[str, Any], facts_by_source: dict[str, list[dict[str, Any]]], source_map: dict[str, dict[str, Any]]) -> tuple[str, str]:
    if text(lineage.get('Table Evidence')):
        return table_match(lineage, source_map)
    candidates = matching_facts(lineage, facts_by_source, source_map)
    expected = dec(lineage.get("Value"))
    if not candidates:
        return "NOT_RUN", "Evidence不足: no exact concept/entity/period/unit/dimension/face-scope candidate."
    intervals = [rounding_interval(fact) for fact in candidates]
    if max(low for low, _ in intervals) > min(high for _, high in intervals):
        return "NOT_RUN", "Evidence不足: inconsistent duplicate facts with identical semantic aspects."
    matches = [fact for fact in candidates if fact["value"] == expected]
    if not matches:
        return "FAIL", f"金额错误: selected {expected}; same-scope SEC facts {sorted({str(f['value']) for f in candidates})}."
    fact = matches[0]
    return "PASS", f"SEC inline fact {fact['id']} at {fact['xpath']}; context={fact['context']}; unit={fact['unit']}; raw={fact['raw_lexical']}; scale={fact['scale']}; sign={fact['sign']}; artifact={fact['artifact_sha256']}."


def table_match(lineage: dict[str, Any], source_map: dict[str, dict[str, Any]]) -> tuple[str, str]:
    """Reopen the original artifact and independently inspect the exact column."""
    try:
        evidence = json.loads(lineage['Table Evidence'])
        source = source_map[text(lineage['Source ID'])]
        path = xml_cache_for(text(source['Accession']))
        if path is None:
            raise ValueError('Missing unique SEC artifact')
        stat = path.stat()
        raw, root = parsed_document(path, stat.st_mtime_ns, stat.st_size)
        if hashlib.sha256(raw).hexdigest() != evidence['artifact_sha256']:
            raise ValueError('TABLE artifact hash differs')
        cells = root.xpath(evidence['xpath'])
        if len(cells) != 1:
            raise ValueError('TABLE locator not unique')
        cell = cells[0]
        table = cell.xpath('ancestor::*[local-name()="table"][1]')[0]
        unit_caption = ' '.join(table.itertext()) + ' '.join(' '.join(node.itertext()) for node in table.xpath('preceding-sibling::*[position() <= 4]'))
        wrapper = table.getparent()
        if wrapper is not None and len(wrapper.xpath('./*[local-name()="table"]')) == 1:
            unit_caption += ' '.join(' '.join(node.itertext()) for node in wrapper.xpath('preceding-sibling::*[position() <= 5]'))
        if 'in millions' not in unit_caption.lower() or Decimal(evidence['scale']) != Decimal('1000000'):
            raise ValueError('TABLE local disclosure scale not established')
        tr = cell.xpath('ancestor::*[local-name()="tr"][1]')[0]
        label = ''.join(tr.xpath('./*[local-name()="td" or local-name()="th"][1]')[0].itertext())
        if aspect_name(label) != aspect_name(evidence['label']):
            raise ValueError('TABLE row label differs')
        if evidence['year'] not in ' '.join(table.itertext()):
            raise ValueError('TABLE target year missing')
        # Independent span expansion for the target header, not a call into
        # the production locator or a comparison against its saved raw value.
        rows = table.xpath('./*[local-name()="tr"] | ./*[local-name()="tbody" or local-name()="thead"]/*[local-name()="tr"]')
        occupied = {}
        target_column = None
        preceding = []
        for ri, rr in enumerate(rows):
            col = 0
            for cc in rr.xpath('./*[local-name()="td" or local-name()="th"]'):
                while (ri, col) in occupied:
                    col += 1
                for dy in range(int(cc.get('rowspan', '1'))):
                    for dx in range(int(cc.get('colspan', '1'))):
                        occupied[ri + dy, col + dx] = cc
                if cc is cell:
                    target_column = col
                col += int(cc.get('colspan', '1'))
        if target_column is None:
            raise ValueError('TABLE physical cell missing')
        target_row = rows.index(tr)
        preceding = [' '.join(occupied[ri, target_column].itertext()).strip() for ri in range(target_row) if (ri, target_column) in occupied]
        if evidence['year'] not in preceding or aspect_name(evidence['header']) not in aspect_name(' '.join(preceding)):
            raise ValueError('TABLE target column/header differs')
        period = parse_period(text(lineage['Context']))
        if period != ('duration', evidence['start'], evidence['end']) or lineage['Unit'] != evidence['unit']:
            raise ValueError('TABLE period/unit differs')
        value = dec(''.join(cell.itertext()))
        if value is None or value * Decimal(evidence['scale']) != dec(lineage['Value']):
            return 'FAIL', '金额错误: TABLE lexical amount differs from canonical output.'
        return 'PASS', f"SEC TABLE row={label.strip()}; column={target_column}; header={evidence['header']} {evidence['year']}; raw={value}; locator={evidence['xpath']}; artifact={evidence['artifact_sha256']}."
    except (KeyError, ValueError, IndexError, TypeError, json.JSONDecodeError) as error:
        return 'NOT_RUN', 'Evidence不足: TABLE verification: ' + str(error)


def verify_evidence(lineage: dict[str, Any], facts_by_source: dict[str, list[dict[str, Any]]], source_map: dict[str, dict[str, Any]], depth: int = 0) -> tuple[str, str]:
    if depth > 12:
        return "NOT_RUN", "Evidence不足: recursive inputs exceed depth limit."
    if text(lineage.get("Type")) not in {"C", "D"}:
        return xbrl_match(lineage, facts_by_source, source_map)
    raw = text(lineage.get("Input Facts"))
    try:
        inputs = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return "NOT_RUN", "Evidence不足: no structured derivation inputs."
    if not isinstance(inputs, list) or not inputs:
        return "NOT_RUN", "Evidence不足: empty derivation inputs."
    total = Decimal(0)
    dependency_notes = []
    intervals = []
    accession_map = {text(source["Accession"]): sid for sid, source in source_map.items()}
    for item in inputs:
        amount, coefficient = dec(item.get("Value")), dec(item.get("Coefficient"))
        if amount is None or coefficient is None:
            return "NOT_RUN", "Evidence不足: invalid input value/coefficient."
        start, end = text(item.get("Start")), text(item.get("End"))
        nested = text(item.get("SourceInputFacts"))
        if nested.casefold() in {'nan', 'none', 'null'}:
            nested = ''
        child = {"Source ID": accession_map.get(text(item.get("Accession")), ""),
                 "XBRL Concept": item.get("Concept"), "Value": amount,
                 "Context": (f"{start}..{end}" if start and start != end else end) + (
                     ' | ' + text(item.get('SourceEvidenceAxes')) + '; ' + text(item.get('SourceEvidenceMembers'))
                     if text(item.get('SourceEvidenceAxes')) not in {'', 'nan', 'None'} else ''),
                 "Metric": item.get("Label", lineage.get("Metric")),
                 "Unit": item.get("Unit") or lineage.get("Unit"),
                 "Type": ("D" if item.get("Concept") else "C") if nested and nested not in {"nan", "none"} else "R", "Input Facts": nested,
                 'Table Evidence': item.get('SourceTableEvidence') if text(item.get('SourceTableEvidence')) not in {'nan', 'None'} else ''}
        result, diagnostic = verify_evidence(child, facts_by_source, source_map, depth + 1)
        if result != "PASS":
            return result, "Input: " + diagnostic
        total += amount * coefficient
        dependency_notes.append(diagnostic)
        intervals.append((start, end, coefficient))
    if total != dec(lineage.get("Value")):
        return "FAIL", f"金额错误: independently recomputed inputs={total}, output={lineage.get('Value')}."
    if text(lineage.get("Type")) == "D":
        local_concept = text(lineage.get("XBRL Concept")).split('}')[-1].split(':')[-1]
        if local_concept.startswith(("EarningsPerShare", "WeightedAverage")):
            return "FAIL", "口径错误: nonadditive metric cannot be subtracted."
        target = parse_period(text(lineage.get("Context")))
        if not target or target[0] != "duration":
            return "NOT_RUN", "Evidence不足: no duration coverage proof."
        points = sorted({target[1], target[2], *(date for start, end, _ in intervals for date in (start, end))})
        # SEC duration ends are inclusive: use explicit boundary dates.
        from datetime import date, timedelta
        segments = [(date.fromisoformat(start), date.fromisoformat(end) + timedelta(days=1), coef) for start, end, coef in intervals]
        target_start, target_end = date.fromisoformat(target[1]), date.fromisoformat(target[2]) + timedelta(days=1)
        boundaries = sorted({target_start, target_end, *(p for start, end, _ in segments for p in (start, end))})
        for left, right in zip(boundaries, boundaries[1:]):
            coverage = sum(coef for start, end, coef in segments if start <= left and end >= right)
            if coverage != (1 if target_start <= left and right <= target_end else 0):
                return "FAIL", "口径错误: cumulative intervals do not cover the output quarter exactly."
        del points
    return "PASS", "Independent Decimal input recomputation and period proof. " + " || ".join(dependency_notes)


def front_values(book: Any, lineage_map: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for sheet in MAIN_SHEETS:
        ws = book[sheet]
        for row in range(1, ws.max_row + 1):
            label = ws.cell(row, 2).value
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row, col)
                if not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool):
                    continue
                location = text(cell.hyperlink.location) if cell.hyperlink else ""
                match = re.fullmatch(r"'91_LINEAGE'!A(\d+)", location)
                lineage = None
                if match:
                    lineage = lineage_map.get(f"row:{match.group(1)}")
                rows.append({
                    "Sheet": sheet, "Cell": cell.coordinate, "Label": text(lineage["Metric"]) if lineage else text(label),
                    "Period": text(lineage["Period"]) if lineage else (text(ws.cell(5, col).value) if row >= 5 else ""),
                    "Displayed Value": cell.value, "Binding": location,
                    "Lineage ID": lineage["Lineage ID"] if lineage else "",
                    "Scope": "metadata" if sheet == "00_OVERVIEW" and cell.coordinate in {"C10", "C11", "C12"} else "financial",
                })
    return rows


def evaluate_front(rows: list[dict[str, Any]], lineage_by_row: dict[str, dict[str, Any]], facts_by_source: dict[str, list[dict[str, Any]]], source_map: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for item in rows:
        result = dict(item)
        # Overview counts are numeric cells, but they are workbook metadata,
        # not financial outputs.  Keep them in the audit population while
        # excluding them from the financial-value conclusion.
        if item.get("Scope") == "metadata":
            result.update({
                "Result": "NOT_APPLICABLE",
                "Value Check": "NOT_APPLICABLE",
                "Evidence Check": "NOT_APPLICABLE",
                "Diagnostic": "Workbook metadata, excluded from financial numeric denominator.",
            })
            results.append(result)
            continue
        lid = text(item["Lineage ID"])
        lineage = lineage_by_row.get(lid)
        if lineage is None:
            result.update({"Result": "FAIL", "Value Check": "FAIL", "Evidence Check": "FAIL", "Diagnostic": "Numeric front cell has no valid 91_LINEAGE binding."})
            results.append(result)
            continue
        expected = dec(lineage["Value"])
        scale = dec(lineage["Display Scale"]) or Decimal("1")
        display_ok = expected is not None and close(item["Displayed Value"], expected * scale)
        result["Lineage Value"] = text(lineage["Value"])
        result["Type"] = text(lineage["Type"])
        result["Source ID"] = text(lineage["Source ID"])
        result["XBRL Concept"] = text(lineage["XBRL Concept"])
        result["Context"] = text(lineage["Context"])
        result["Unit"] = text(lineage["Unit"])
        result["Display Scale"] = text(lineage["Display Scale"])
        result["Locator"] = text(lineage["Locator"])
        result["Value Check"] = "PASS" if display_ok else "FAIL"
        if not display_ok:
            result.update({"Result": "FAIL", "Evidence Check": "FAIL", "Diagnostic": "Front display does not equal lineage value times display scale."})
        elif not text(lineage["Source ID"]) or text(lineage["Source ID"]) not in source_map:
            result.update({"Result": "FAIL", "Evidence Check": "FAIL", "Diagnostic": "Missing or unknown Source ID."})
        elif text(lineage["Type"]) == "R" and (not text(lineage["XBRL Concept"]) and not text(lineage["Locator"])):
            result.update({"Result": "NOT_RUN", "Evidence Check": "NOT_RUN", "Diagnostic": "Reported value has neither concept nor locator."})
        elif text(lineage["Type"]) in {"C", "D"} and not text(lineage["Formula"]):
            result.update({"Result": "NOT_RUN", "Evidence Check": "NOT_RUN", "Diagnostic": "Derived/calculated value has no formula."})
        else:
            evidence_result, diagnostic = verify_evidence(lineage, facts_by_source, source_map)
            result.update({"Result": evidence_result, "Evidence Check": evidence_result, "Diagnostic": diagnostic})
        results.append(result)
    return results


def parse_input_facts(lineage: dict[str, Any]) -> tuple[Decimal | None, str]:
    raw = text(lineage["Input Facts"])
    if not raw or raw.lower() in {"nan", "none"}:
        return None, "No structured input facts."
    try:
        inputs = json.loads(raw)
    except json.JSONDecodeError:
        return None, "Input Facts is not valid JSON."
    total = Decimal("0")
    for item in inputs:
        value, coef = dec(item.get("Value")), dec(item.get("Coefficient"))
        if value is None or coef is None:
            return None, "Input fact has invalid value or coefficient."
        total += value * coef
    return total, "Structured inputs recomputed with Decimal."


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def fault_injection(workbook: Path) -> list[dict[str, Any]]:
    from excel_lineage import verify_workbook_manifest
    baseline = verify_workbook_manifest(workbook)
    if baseline:
        return [{"Case": "baseline", "Rejected": False, "Diagnostics": "Baseline invalid: " + " | ".join(baseline)}]
    cases = ["front_value_changed", "wrong_lineage_binding", "display_scale_changed", "unit_changed", "source_removed", "qa_removed", "input_removed"]
    results = []
    with tempfile.TemporaryDirectory(prefix="google_review_") as directory:
        for case in cases:
            target = Path(directory) / (case + ".xlsx")
            shutil.copy2(workbook, target)
            manifest_path = target.with_suffix(".manifest.json")
            manifest = json.loads(workbook.with_suffix(".manifest.json").read_text(encoding="utf-8"))
            w = load_workbook(target)
            front = w["02_QUARTERLY"]["F6"]
            target_row = int(text(front.hyperlink.location).split("A")[-1])
            audit = w["91_LINEAGE"]
            headers = {text(audit.cell(4, col).value): col for col in range(2, audit.max_column + 1)}
            if case == "front_value_changed":
                front.value += 1
            elif case == "wrong_lineage_binding":
                from openpyxl.worksheet.hyperlink import Hyperlink
                front.hyperlink = Hyperlink(ref=front.coordinate, location=f"'91_LINEAGE'!A{target_row + 1}")
            elif case in {"display_scale_changed", "unit_changed"}:
                field = "Display Scale" if case == "display_scale_changed" else "Unit"
                audit.cell(target_row, headers[field]).value = 2e-6 if field == "Display Scale" else "shares"
            elif case == "source_removed":
                w["90_SOURCES"]["B5"] = ""
            elif case == "qa_removed":
                w["08_QA"]["B5"] = ""
            else:
                row = next(row for row in range(5, audit.max_row + 1) if audit.cell(row, headers["Input Facts"]).value)
                audit.cell(row, headers["Input Facts"]).value = None
            w.save(target)
            w.close()
            manifest["workbook"].update(sha256=sha256(target), bytes=target.stat().st_size)
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            errors = verify_workbook_manifest(target)
            results.append({"Case": case, "Baseline": "PASS", "Binary digest refreshed": True, "Rejected": bool(errors), "Diagnostics": " | ".join(errors)})
            print(json.dumps({'phase': 'fault_injection', 'case': case, 'rejected': bool(errors)}), flush=True)
    return results


def main() -> None:
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=OUT)
    args = parser.parse_args()
    output = args.output_dir
    output.mkdir(parents=True, exist_ok=True)
    book = load_workbook(args.workbook)
    sources, source_map = source_rows(book)
    lineages, lineage_map = lineage_rows(book)
    by_row = {f"row:{row}": lineage_map.get(text(book["91_LINEAGE"].cell(row, 2).value), {}) for row in range(5, book["91_LINEAGE"].max_row + 1)}
    front = front_values(book, by_row)
    facts_by_source = {}
    inventory = []
    for sid, source in source_map.items():
        path = xml_cache_for(text(source["Accession"]))
        if path is None:
            inventory.append({"Source ID": sid, "Status": "NOT_RUN", "Reason": "Missing or ambiguous cached SEC bytes"})
            continue
        _, facts = xml_facts(path)
        facts_by_source[sid] = facts
        inventory.append({"Source ID": sid, "Status": "PASS", "Cache File": path.name, "SHA256": sha256(path), "Facts": len(facts)})
    results = evaluate_front(front, lineage_map, facts_by_source, source_map)
    print(json.dumps({'phase': 'numeric_review', 'results': dict(Counter(row['Result'] for row in results))}), flush=True)
    for row in results:
        row["Issue Category"] = ("金额错误" if "金额错误" in row.get("Diagnostic", "") else "口径错误" if "口径错误" in row.get("Diagnostic", "") else "证据不足" if row["Result"] == "NOT_RUN" else "核验工具问题" if row["Result"] == "FAIL" else "")
    faults = fault_injection(args.workbook)
    from excel_lineage import verify_workbook_manifest
    checks = [{"Check": "production_integrity", "Result": "FAIL" if verify_workbook_manifest(args.workbook) else "PASS"},
              {"Check": "network_smoke", "Result": "NOT_RUN", "Reason": "Offline cached evidence review"}]
    # Independent cash bridge: the engine's before-FX subtotal is not itself
    # evidence for the change in the cash universe. Use reported instant facts.
    from datetime import date, timedelta
    cash_facts = [fact for facts in facts_by_source.values() for fact in facts
                  if fact['concept'] == 'CashAndCashEquivalentsAtCarryingValue'
                  and fact['kind'] == 'instant' and fact['unit'] == 'USD' and not fact['dimensions']]
    cash_rows = {(text(row['Metric']), text(row['Period'])): row for row in lineages if row['Excel Sheet'] == '05_CASH_FLOW'}
    displayed_cash_periods = {text(item['Period']) for item in front if item['Label'] == 'Operating Cash Flow' and item['Scope'] == 'financial'}
    for row in lineages:
        if row['Excel Sheet'] != '05_CASH_FLOW' or row['Metric'] != 'Operating Cash Flow':
            continue
        if text(row['Period']) not in displayed_cash_periods:
            continue
        period = parse_period(text(row['Context']))
        if not period or period[0] != 'duration':
            continue
        beginning = (date.fromisoformat(period[1]) - timedelta(days=1)).isoformat()
        end = period[2]
        operands = [cash_rows.get((label, text(row['Period']))) for label in
                    ['Operating Cash Flow', 'Investing Cash Flow', 'Financing Cash Flow', 'FX Effect on Cash']]
        starts = [fact for fact in cash_facts if fact['start'] == beginning]
        ends = [fact for fact in cash_facts if fact['start'] == end]
        check = {'Check': 'independent_cash_bridge', 'Period': row['Period'], 'Beginning': beginning, 'Ending': end}
        if any(item is None or dec(item.get('Value')) is None for item in operands) or len({fact['value'] for fact in starts}) != 1 or len({fact['value'] for fact in ends}) != 1:
            check.update(Result='NOT_RUN', Reason='Missing or conflicting cash-universe instant/flow inputs')
        else:
            difference = starts[0]['value'] + sum(dec(item['Value']) for item in operands) - ends[0]['value']
            # Each disclosed USD-million item has a half-million rounding
            # interval; propagation uses the actual reported decimals.
            raw_inputs = starts[:1] + ends[:1]
            tolerance = sum((rounding_interval(fact)[1] - fact['value'] for fact in raw_inputs), Decimal(0))
            flow_intervals = []
            for item in operands:
                matches = matching_facts(item, facts_by_source, source_map)
                if matches:
                    flow_intervals.append(rounding_interval(matches[0])[1] - matches[0]['value'])
            if len(flow_intervals) != 4 and difference != 0:
                check.update(Result='NOT_RUN', Reason='Derived-flow disclosure rounding not independently bounded', Difference=str(difference))
            else:
                tolerance += sum(flow_intervals, Decimal(0))
                check.update(Result='PASS' if abs(difference) <= tolerance else 'FAIL', Difference=str(difference),
                             Tolerance=str(tolerance), Method='Reported beginning cash + CFO + CFI + CFF + FX = reported ending cash')
        checks.append(check)
    write_csv(output / "front_value_review.csv", results)
    write_csv(output / "source_cache_inventory.csv", inventory)
    write_csv(output / "fault_injection.csv", faults)
    write_csv(output / "structural_checks.csv", checks)
    write_csv(output / "defects.csv", [row for row in results if row["Result"] in {"FAIL", "NOT_RUN"}])
    summary = {"version": "luna-corrected-v2", "workbook": args.workbook.name, "sha256": sha256(args.workbook),
               "executed_at": datetime.now(timezone.utc).isoformat(), "front_financial_cells": sum(row["Scope"] == "financial" for row in results),
               "unique_data_points": len({row["Lineage ID"] for row in results if row["Scope"] == "financial"}),
               "results": dict(Counter(row["Result"] for row in results)), "faults_rejected": sum(row["Rejected"] for row in faults),
               "conclusion": "PASS" if all(row["Result"] in {"PASS", "NOT_APPLICABLE"} for row in results) and all(row["Rejected"] for row in faults) else "REWORK"}
    (output / "review_manifest.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    book.close()
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
