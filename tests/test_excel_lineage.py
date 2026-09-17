import hashlib
import json
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from excel_lineage import (
    SHEET_NAMES,
    _program_qa,
    _source_rows,
    save_pivot_xlsx,
    verify_workbook_manifest,
)


def sample_pivot() -> pd.DataFrame:
    index = pd.MultiIndex.from_tuples(
        [
            ("0_Period_Header", "Period Ending"),
            ("1_Income_Statement", "Revenue"),
            ("3_Cash_Flow", "Operating Cash Flow"),
            ("5_KPI_Metrics", "Adjusted EBITDA"),
        ],
        names=["Category", "Label"],
    )
    frame = pd.DataFrame(
        {"2026-Q1": ["2026-03-31", 487_000_000, 226_000_000, 426_000_000]},
        index=index,
    )
    frame.attrs["fact_audit"] = pd.DataFrame(
        [
            {
                "Ticker": "TLN", "Category": "1_Income_Statement", "Label": "Revenue",
                "Period": "2026-Q1", "Value": 487_000_000, "SourceKind": "xbrl",
                "SourceAdmissionRule": "consolidated_concept_map", "Concept": "Revenues",
                "Start": "2026-01-01", "End": "2026-03-31", "Form": "10-Q",
                "Filed": "2026-05-05", "Accession": "0001622536-26-000036",
                "FilingUrl": "https://www.sec.gov/example.htm",
                "AcceptedAt": "2026-05-05T16:01:02-04:00",
            },
            {
                "Ticker": "TLN", "Category": "3_Cash_Flow", "Label": "Operating Cash Flow",
                "Period": "2026-Q1", "Value": 226_000_000, "SourceKind": "xbrl",
                "SourceAdmissionRule": "consolidated_concept_map",
                "SourceDerivation": "exact_concept_ytd6_minus_q1",
                "SourceDerivationFormula": "H1 YTD - Q1 YTD",
                "Concept": "NetCashProvidedByUsedInOperatingActivities",
                "Start": "2026-01-01", "End": "2026-03-31", "Form": "10-Q",
                "Filed": "2026-05-05", "Accession": "0001622536-26-000036",
                "FilingUrl": "https://www.sec.gov/example.htm",
            },
            {
                "Ticker": "TLN", "Category": "5_KPI_Metrics", "Label": "Adjusted EBITDA",
                "Period": "2026-Q1", "Value": 426_000_000, "SourceKind": "calculated",
                "SourceAdmissionRule": "kpi_calculation", "Concept": "",
            },
        ]
    )
    return frame


class ExcelLineageTests(unittest.TestCase):
    def _build(self, frame: pd.DataFrame | None = None):
        temp = tempfile.TemporaryDirectory()
        path = Path(temp.name) / "test.xlsx"
        save_pivot_xlsx(frame if frame is not None else sample_pivot(), path)
        workbook = load_workbook(path, data_only=False)
        return temp, path, workbook

    def test_sheet_order_and_clean_front(self):
        temp, _, workbook = self._build()
        self.addCleanup(temp.cleanup)
        self.assertEqual(tuple(workbook.sheetnames), SHEET_NAMES)
        for sheet in ("03_INCOME_STATEMENT", "04_BALANCE_SHEET", "05_CASH_FLOW"):
            values = {cell.value for row in workbook[sheet].iter_rows() for cell in row}
            self.assertNotIn("Accession", values)
            self.assertNotIn("XBRL Concept", values)
            self.assertNotIn("Source ID", values)

    def test_overview_exposes_layered_coverage(self):
        frame = sample_pivot()
        frame.attrs["fact_audit"]["SourceClassificationConfidence"] = 0.99
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        overview = workbook["00_OVERVIEW"]
        rows = {
            overview.cell(row, 2).value: overview.cell(row, 3).value
            for row in range(1, overview.max_row + 1)
        }
        self.assertEqual(rows["Financial Core"], "REVIEW")
        self.assertEqual(rows["Industry Mapping"], "POWER / SUPPORTED")
        self.assertEqual(rows["Company KPI"], "CONFIGURED")

    def test_lineage_exposes_semantic_confidence(self):
        frame = sample_pivot()
        frame.attrs["fact_audit"].loc[0, "SourceSemanticType"] = "Revenue"
        frame.attrs["fact_audit"].loc[0, "SourceClassificationConfidence"] = 0.99
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        lineage = workbook["91_LINEAGE"]
        headers = {
            lineage.cell(4, col).value: col
            for col in range(2, lineage.max_column + 1)
        }
        row = next(
            row for row in range(5, lineage.max_row + 1)
            if lineage.cell(row, headers["Metric"]).value == "Revenue"
        )
        self.assertEqual(
            lineage.cell(row, headers["Semantic Classification"]).value,
            "Revenue",
        )
        self.assertEqual(
            lineage.cell(row, headers["Semantic Confidence"]).value,
            "HIGH",
        )
        self.assertEqual(
            lineage.cell(row, headers["Review Required"]).value,
            "NO",
        )

    def test_source_is_stored_once_and_referenced(self):
        temp, _, workbook = self._build()
        self.addCleanup(temp.cleanup)
        sources = workbook["90_SOURCES"]
        accessions = [sources.cell(row, 8).value for row in range(5, sources.max_row + 1)]
        self.assertEqual(accessions, ["0001622536-26-000036"])
        lineage = workbook["91_LINEAGE"]
        headers = {lineage.cell(4, col).value: col for col in range(2, lineage.max_column + 1)}
        source_ids = [lineage.cell(row, headers["Source ID"]).value for row in range(5, lineage.max_row + 1)]
        self.assertGreaterEqual(source_ids.count("F001"), 2)
        link_cell = sources.cell(5, 9)
        self.assertEqual(link_cell.value, "SEC 原文 ↗")
        self.assertEqual(link_cell.hyperlink.target, "https://www.sec.gov/example.htm")
        self.assertEqual(link_cell.font.color.rgb, "00374151")
        self.assertIsNone(link_cell.font.underline)

    def test_source_period_uses_filing_end_not_first_comparative_row(self):
        audit = sample_pivot().attrs["fact_audit"].copy()
        audit["FY"] = 2026
        audit["Q"] = "Q1"
        comparative = audit.iloc[0].copy()
        comparative["FY"] = 2025
        comparative["Q"] = "Q4"
        comparative["Start"] = "2025-01-01"
        comparative["End"] = "2025-12-31"
        audit = pd.concat([pd.DataFrame([comparative]), audit], ignore_index=True)
        rows, _ = _source_rows(audit)
        self.assertEqual(rows[0]["Report Period"], "2026 Q1")

    def test_exact_cell_link_and_origin(self):
        temp, _, workbook = self._build()
        self.addCleanup(temp.cleanup)
        income = workbook["03_INCOME_STATEMENT"]
        revenue_cell = next(cell for row in income.iter_rows() for cell in row if cell.value == 487)
        self.assertIsNotNone(revenue_cell.hyperlink)
        self.assertTrue(revenue_cell.hyperlink.location.startswith("'91_LINEAGE'!A"))
        lineage = workbook["91_LINEAGE"]
        headers = {lineage.cell(4, col).value: col for col in range(2, lineage.max_column + 1)}
        rows = {
            lineage.cell(row, headers["Metric"]).value: row
            for row in range(5, lineage.max_row + 1)
        }
        self.assertEqual(lineage.cell(rows["Revenue"], headers["Type"]).value, "R")
        self.assertEqual(lineage.cell(rows["Operating Cash Flow"], headers["Type"]).value, "D")
        self.assertEqual(lineage.cell(rows["Operating Cash Flow"], headers["Formula"]).value, "H1 YTD - Q1 YTD")

    def test_missing_audit_is_program_qa_not_ai_review(self):
        frame = sample_pivot()
        frame.attrs["fact_audit"] = frame.attrs["fact_audit"].iloc[:2].copy()
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        raw_qa = workbook["98_RAW_QA"]
        findings = [raw_qa.cell(row, 9).value for row in range(5, raw_qa.max_row + 1)]
        self.assertTrue(any("period-level" in str(value) for value in findings))
        review = workbook["09_REVIEW"]
        review_findings = [review.cell(row, 11).value for row in range(5, review.max_row + 1)]
        self.assertFalse(any("No matching fact_audit" in str(value) for value in review_findings))
        lineage = workbook["91_LINEAGE"]
        headers = {lineage.cell(4, col).value: col for col in range(2, lineage.max_column + 1)}
        kpi_row = next(
            row for row in range(5, lineage.max_row + 1)
            if lineage.cell(row, headers["Metric"]).value == "Adjusted EBITDA"
        )
        self.assertEqual(lineage.cell(kpi_row, headers["Type"]).value, "C")
        self.assertTrue(lineage.cell(kpi_row, headers["Source ID"]).value)
        self.assertIn(
            "period-level input filings",
            lineage.cell(kpi_row, headers["Method"]).value,
        )

    def test_source_text_cannot_become_formula(self):
        frame = sample_pivot()
        frame.attrs["fact_audit"].loc[0, "Concept"] = "=HYPERLINK(\"https://bad\",\"x\")"
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        lineage = workbook["91_LINEAGE"]
        headers = {lineage.cell(4, col).value: col for col in range(2, lineage.max_column + 1)}
        cell = lineage.cell(5, headers["XBRL Concept"])
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(str(cell.value).startswith("'="))

    def test_backend_sheets_hidden_and_review_queue_bounded(self):
        temp, _, workbook = self._build()
        self.addCleanup(temp.cleanup)
        for name in ("90_SOURCES", "91_LINEAGE", "97_RAW_KPI", "98_RAW_QA"):
            self.assertEqual(workbook[name].sheet_state, "hidden")
        review = workbook["09_REVIEW"]
        self.assertLessEqual(max(review.max_row - 4, 0), 30)

    def test_annual_companion_populates_annual_summary(self):
        annual = sample_pivot().rename(columns={"2026-Q1": "2025-FY"})
        annual.attrs["fact_audit"] = annual.attrs["fact_audit"].assign(Period="2025-FY", Form="10-K")
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / "combined.xlsx"
        save_pivot_xlsx(sample_pivot(), path, annual_pivot=annual)
        workbook = load_workbook(path, data_only=False)
        values = {cell.value for row in workbook["01_ANNUAL"].iter_rows() for cell in row}
        self.assertIn("2025-FY", values)
        self.assertIn(487, values)

    def test_resolved_per_mwh_scale_is_not_sent_to_ai_review(self):
        frame = sample_pivot()
        frame.loc[("4a_Segments_Business", "Other Segment Items - $/MWh - PJM West Hub ATC"), :] = [71_080_000]
        extra = frame.attrs["fact_audit"].iloc[0].to_dict()
        extra.update({
            "Category": "4a_Segments_Business",
            "Label": "Other Segment Items - $/MWh - PJM West Hub ATC",
            "Value": 71_080_000,
            "SourceRawValue": "71.08",
            "SourceUnitKind": "money",
            "SourceUnitScale": 1_000_000,
            "SourceSemanticType": "PJM West Hub ATC power price",
            "SourceClassificationConfidence": 0.99,
        })
        frame.attrs["fact_audit"] = pd.concat([frame.attrs["fact_audit"], pd.DataFrame([extra])], ignore_index=True)
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        review = workbook["09_REVIEW"]
        findings = [review.cell(row, 11).value for row in range(5, review.max_row + 1)]
        self.assertFalse(any("unit scaling" in str(value).casefold() for value in findings))
        kpi = workbook["07_OPERATING_KPI"]
        displayed = {cell.value for row in kpi.iter_rows() for cell in row}
        self.assertIn(71.08, displayed)

    def test_review_does_not_compare_nonconsecutive_quarters(self):
        frame = sample_pivot().copy()
        frame["2016-Q3"] = frame["2026-Q1"]
        frame["2024-Q1"] = frame["2026-Q1"] * 10
        frame = frame.drop(columns=["2026-Q1"])
        audit = sample_pivot().attrs["fact_audit"]
        old = audit.assign(Period="2016-Q3", FY=2016, Q="Q3", End="2016-09-30")
        recent = audit.assign(Period="2024-Q1", FY=2024, Q="Q1", End="2024-03-31")
        frame.attrs["fact_audit"] = pd.concat([old, recent], ignore_index=True)
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        review = workbook["09_REVIEW"]
        findings = [review.cell(row, 11).value for row in range(5, review.max_row + 1)]
        self.assertFalse(any("2016-Q3" in str(value) for value in findings))

    def test_low_confidence_calculation_is_not_presented_as_trusted(self):
        frame = sample_pivot()
        frame.loc[("5_KPI_Metrics", "Metric: Total Debt"), :] = [57_195_000]
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        overview = workbook["00_OVERVIEW"]
        total_debt_row = next(
            row for row in range(1, overview.max_row + 1)
            if str(overview.cell(row, 5).value).startswith("Total Debt")
        )
        self.assertEqual(overview.cell(total_debt_row, 6).value, "—")
        review = workbook["09_REVIEW"]
        self.assertTrue(any(
            review.cell(row, 5).value == "Metric: Total Debt"
            for row in range(5, review.max_row + 1)
        ))

    def test_engine_residual_is_auditable_but_not_a_face_statement_row(self):
        frame = sample_pivot()
        label = 'Other Operating Adjustments (Net)'
        frame.loc[('3_Cash_Flow', label), :] = [123]
        temp, _, workbook = self._build(frame)
        self.addCleanup(temp.cleanup)
        self.assertNotIn(label, {cell.value for row in workbook['05_CASH_FLOW'] for cell in row})
        self.assertIn(label, {cell.value for row in workbook['91_LINEAGE'] for cell in row})

    def test_basic_accounting_qa_uses_decimal_equations(self):
        def item(metric, value, period="2026-Q1"):
            return {
                "_Category": "2_Balance_Sheet" if metric.startswith("Total") else "3_Cash_Flow",
                "Metric": metric,
                "Period": period,
                "Value": value,
                "Type": "R",
                "Source ID": "F001",
                "XBRL Concept": metric.replace(" ", ""),
                "Unit": "USD",
                "Review Required": "NO",
            }

        lineage = [
            item("Total Assets", "1000000000.1"),
            item("Total Liabilities", "400000000.1"),
            item("Total Equity", "600000000.0"),
            item("Operating Cash Flow", "10.1"),
            item("Investing Cash Flow", "-3.2"),
            item("Financing Cash Flow", "-1.9"),
            item("Net Cash Flow", "5.0"),
        ]
        rows = _program_qa(lineage, [{"Source ID": "F001"}], [])
        statuses = {row["Check ID"]: row["Status"] for row in rows}
        self.assertEqual(statuses["QA006"], "PASS")
        self.assertEqual(statuses["QA007"], "PASS")

    def test_build_manifest_binds_workbook_bytes(self):
        temp, path, _ = self._build()
        self.addCleanup(temp.cleanup)
        self.assertEqual(verify_workbook_manifest(path), [])
        path.write_bytes(path.read_bytes() + b"tampered")
        self.assertEqual(
            verify_workbook_manifest(path),
            ["workbook hash does not match build manifest"],
        )

    def test_manifest_rejects_semantic_change_even_with_refreshed_file_hash(self):
        temp, path, _ = self._build()
        self.addCleanup(temp.cleanup)
        workbook = load_workbook(path)
        lineage = workbook["91_LINEAGE"]
        headers = {
            lineage.cell(4, column).value: column
            for column in range(2, lineage.max_column + 1)
        }
        lineage.cell(5, headers["Value"], 999_999_999)
        workbook.save(path)
        manifest_path = path.with_suffix(".manifest.json")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["workbook"]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        manifest["workbook"]["bytes"] = path.stat().st_size
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        self.assertEqual(
            verify_workbook_manifest(path),
            ["workbook semantic content does not match build manifest"],
        )

    def test_front_tampering_is_rejected_with_refreshed_file_hash(self):
        temp, path, book = self._build()
        self.addCleanup(temp.cleanup)
        sheet = book['03_INCOME_STATEMENT']
        row = next(r for r in range(7, sheet.max_row + 1) if sheet.cell(r, 2).value == 'Revenue')
        sheet.cell(row, 3, 999)
        book.save(path)
        manifest_path = path.with_suffix('.manifest.json')
        manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        manifest['workbook']['sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
        manifest_path.write_text(json.dumps(manifest), encoding='utf-8')
        self.assertEqual(verify_workbook_manifest(path),
                         [f'03_INCOME_STATEMENT!C{row} differs from its lineage value'])


if __name__ == "__main__":
    unittest.main()


def test_period_dates_fall_back_to_unique_audit_endpoints():
    from excel_lineage import _periods
    frame = pd.DataFrame({"2026-Q1": [1]}, index=pd.MultiIndex.from_tuples([("1_Income_Statement", "Revenue")]))
    frame.attrs["fact_audit"] = pd.DataFrame([{"Period": "2026-Q1", "End": "2026-03-31"}])
    periods, dates = _periods(frame)
    assert periods == ["2026-Q1"]
    assert dates["2026-Q1"] == "2026-03-31"


def test_optional_withheld_candidate_does_not_block_minimum_financial_core():
    from excel_lineage import _program_qa
    base = {"_Category": "1_Income_Statement", "Period": "2026-Q1", "Unit": "USD", "Source ID": "F1", "Type": "R", "Review Required": "NO", "XBRL Concept": "Revenue"}
    rows = [dict(base, Metric=name, Value=1) for name in ["Revenue", "Operating Income", "Net Income"]]
    rows.append(dict(base, _Category="2_Balance_Sheet", Metric="Short-term Borrowings", Value=0,
                     **{"Review Required": "YES", "Method": "MISSING_RELOCATABLE_REPORTED_FACT"}))
    qa = {r["Check ID"]: r for r in _program_qa(rows, [{"Source ID": "F1"}], [], {"2026-Q1"})}
    assert qa["QA005"]["Status"] == "PASS"


def test_deterministic_withhold_is_not_sent_for_manual_ai_review():
    from excel_lineage import _semantic_review_queue
    row = {"_Category": "2_Balance_Sheet", "Metric": "Short-term Borrowings", "Period": "2026-Q1",
           "Value": 0, "Review Required": "YES", "Method": "MISSING_RELOCATABLE_REPORTED_FACT",
           "Display Scale": 1, "Unit": "USD"}
    assert _semantic_review_queue([row], {"2026-Q1"}, {("2_Balance_Sheet", "Short-term Borrowings")}) == []


def test_known_engine_rollups_stay_in_raw_qa_not_manual_review():
    from excel_lineage import _semantic_review_queue
    row = {"_Category": "3_Cash_Flow", "Metric": "Other Investing Adjustments (Net)", "Period": "2026-Q1",
           "Value": 1, "Review Required": "YES", "Method": "engine calculation; period-level input filings",
           "Display Scale": 1, "Unit": "USD"}
    assert _semantic_review_queue([row], {"2026-Q1"}, {("3_Cash_Flow", row["Metric"])}) == []
